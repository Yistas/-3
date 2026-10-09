# -*- coding: utf-8 -*-

"""
Диалог для загрузки рыночных цен из Excel файла (оптимизированная версия)
"""

import logging
import re
from datetime import datetime
from pathlib import Path
from collections import defaultdict
from PySide6.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QPushButton,
                               QLabel, QFileDialog, QMessageBox, QProgressBar,
                               QTextEdit, QGroupBox, QGridLayout, QApplication)
from PySide6.QtCore import Qt, QThread, Signal, QTimer
from openpyxl import load_workbook
from database.models import Country, Coin

from utils.paths import paths


class FastLoadPricesThread(QThread):
    """Оптимизированный поток для загрузки цен из Excel"""
    progress = Signal(int, int, str)
    finished = Signal(int, int, list)
    error = Signal(str)

    def __init__(self, db_manager, file_path, file_date):
        super().__init__()
        self.db_manager = db_manager
        self.file_path = file_path
        self.file_date = file_date
        self._is_running = True
        self.logger = logging.getLogger('CoinCollector.LoadPricesThread')

    @staticmethod
    def _parse_price(price_str):
        if not price_str:
            return None
        try:
            cleaned = re.sub(r'[^\d.,]', '', str(price_str))
            if not cleaned:
                return None
            cleaned = cleaned.replace(',', '.')
            return float(cleaned)
        except Exception:
            return None

    @staticmethod
    def _parse_year(year_str):
        if not year_str:
            return None
        try:
            match = re.search(r'(\d{4})', str(year_str))
            return int(match.group(1)) if match else None
        except Exception:
            return None

    @staticmethod
    def _parse_diameter(diameter_str):
        if not diameter_str:
            return None
        try:
            cleaned = re.sub(r'[^\d.,]', '', str(diameter_str))
            if not cleaned:
                return None
            value = float(cleaned)
            return round(value, 1)
        except Exception:
            return None

    @staticmethod
    def _normalize_denomination(value):
        """Строка номинала без пробелов: '12½ сентимо' -> '12½сентимо'"""
        if not value:
            return None
        return re.sub(r'\s+', '', str(value).strip())

    @staticmethod
    def _norm_desc(value):
        """Нормализация описания/наименования для сравнения:
        нижний регистр, без пробелов и знаков препинания."""
        if not value:
            return ''
        t = str(value).lower()
        t = re.sub(r'[^0-9a-zа-яё]+', '', t)
        return t

    @staticmethod
    def _coin_desc_texts(coin):
        """Все текстовые поля монеты, где может лежать наименование/номер."""
        texts = []
        for attr in ('description', 'coin_info', 'obverse_description',
                     'reverse_description', 'purchase_info'):
            v = getattr(coin, attr, None)
            if v:
                texts.append(str(v))
        return texts

    def _find_country(self, country_name, countries_by_name):
        """Ищет страну по названию с учётом альтернативных названий."""
        if not country_name:
            return None
        if country_name in countries_by_name:
            return countries_by_name[country_name]
        alt_names = {
            'Нидерландские Антильские острова': ['Нидерландские Антильские о-ва', 'Антильские о-ва', 'Netherlands Antilles'],
            'Венесуэла': ['Венесуэла (Боливарианская Республика)', 'Боливарианская Республика', 'Venezuela'],
            'Югославия': ['Югославия (Королевство)', 'Королевство Югославия', 'Yugoslavia'],
            'Россия': ['Российская Федерация', 'РФ', 'Russian Federation'],
            'СССР': ['Союз Советских Социалистических Республик', 'Soviet Union', 'USSR'],
            'США': ['Соединенные Штаты Америки', 'United States', 'USA'],
            'Великобритания': ['Соединенное Королевство', 'United Kingdom', 'UK'],
            'Германия': ['ФРГ', 'Германская Федеративная Республика', 'Germany'],
            'Франция': ['Французская Республика', 'France'],
            'Китай': ['КНР', 'Китайская Народная Республика', 'China'],
            'Индия': ['Республика Индия', 'India'],
            'Италия': ['Итальянская Республика', 'Italy'],
            'Испания': ['Королевство Испания', 'Spain'],
            'Япония': ['Государство Япония', 'Japan'],
            'Бразилия': ['Федеративная Республика Бразилия', 'Brazil'],
            'Мексика': ['Мексиканские Соединенные Штаты', 'Mexico'],
            'Канада': ['Canada'],
            'Австралия': ['Содружество Австралии', 'Australia'],
            'ЮАР': ['Южно-Африканская Республика', 'South Africa'],
            'Турция': ['Турецкая Республика', 'Turkey'],
            'Египет': ['Арабская Республика Египет', 'Egypt'],
            'Израиль': ['Государство Израиль', 'Israel'],
            'Иран': ['Исламская Республика Иран', 'Iran'],
            'Ирак': ['Республика Ирак', 'Iraq'],
            'Саудовская Аравия': ['Королевство Саудовская Аравия', 'Saudi Arabia'],
            'ОАЭ': ['Объединенные Арабские Эмираты', 'United Arab Emirates'],
        }
        for main_name, alternatives in alt_names.items():
            if country_name == main_name:
                if main_name in countries_by_name:
                    return countries_by_name[main_name]
            for alt in alternatives:
                if country_name == alt or country_name.lower() == alt.lower():
                    if main_name in countries_by_name:
                        return countries_by_name[main_name]
        for name, country in countries_by_name.items():
            if country_name.lower() in name.lower() or name.lower() in country_name.lower():
                return country
        return None

    def run(self):
        try:
            import tempfile
            from pathlib import Path
            file_path = self.file_path
            file_ext = Path(file_path).suffix.lower()
            temp_file = None
            if file_ext == '.xls':
                try:
                    import pandas as pd
                    df = pd.read_excel(file_path, header=None)
                    temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
                    df.to_excel(temp_file.name, index=False, header=False)
                    file_path = temp_file.name
                except Exception as e:
                    self.error.emit(f"Ошибка конвертации .xls: {e}")
                    return
            wb = load_workbook(file_path, data_only=True, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            if temp_file:
                try:
                    Path(temp_file.name).unlink()
                except Exception:
                    pass
            total_rows = len(rows)
            if total_rows == 0:
                self.error.emit("Файл пуст")
                return
            all_countries = self.db_manager.session.query(Country).all()
            all_coins = self.db_manager.session.query(Coin).all()
            countries_by_name = {}
            for c in all_countries:
                countries_by_name[c.name] = c
            coins_index = defaultdict(list)
            for coin in all_coins:
                denom_normalized = self._normalize_denomination(coin.denomination_value)
                key = (coin.country_id, coin.year, denom_normalized)
                coins_index[key].append(coin)
            updated_count = 0
            error_count = 0
            errors_list = []
            processed_coin_ids = set()
            source_file = Path(self.file_path).name
            for idx, row in enumerate(rows):
                if not self._is_running:
                    break
                if not row or len(row) < 10:
                    continue
                try:
                    country_name = str(row[0]).strip() if row[0] else None
                    if not country_name:
                        continue
                    if country_name.lower() in ["страна", "country", "название", "name"]:
                        continue
                    denomination_raw = str(row[3]).strip() if len(row) > 3 and row[3] else None
                    if denomination_raw and denomination_raw.lower() in ['nan', 'none', '']:
                        denomination_raw = None
                    denomination_normalized = \
                        self._normalize_denomination(denomination_raw) if denomination_raw else None
                    year_str = str(row[4]).strip() if len(row) > 4 and row[4] else None
                    year = self._parse_year(year_str) if year_str else None
                    mint_mark = str(row[5]).strip() if len(row) > 5 and row[5] else None
                    if mint_mark and mint_mark.lower() in ['nan', 'none', '']:
                        mint_mark = None
                    # Колонка 6 — «Наименование»: используется ТОЛЬКО для
                    # сопоставления вариантов монеты, в БД НЕ пишется
                    name_marker = str(row[6]).strip() if len(row) > 6 and row[6] else None
                    if name_marker and name_marker.lower() in ['nan', 'none', '']:
                        name_marker = None
                    diameter_raw = str(row[7]).strip() if len(row) > 7 and row[7] else None
                    diameter = self._parse_diameter(diameter_raw) if diameter_raw else None
                    # КОЛОНКА 10 («Номер», KM#/Y#/UC#) ИГНОРИРУЕТСЯ:
                    # не читается, ни в описание, ни куда-либо ещё НЕ пишется.
                    price_raw = str(row[9]).strip() if len(row) > 9 and row[9] else None
                    if not price_raw or price_raw.lower() in ['nan', 'none', '']:
                        error_count += 1
                        errors_list.append(f"Строка {idx + 2}: нет цены")
                        continue
                    price = self._parse_price(price_raw)
                    if price is None:
                        error_count += 1
                        errors_list.append(
                            f"Строка {idx + 2}: не удалось распарсить цену '{price_raw}'")
                        continue
                    country = countries_by_name.get(country_name)
                    if not country:
                        error_count += 1
                        errors_list.append(f"Строка {idx + 2}: страна '{country_name}' не найдена")
                        continue
                    key = (country.id, year, denomination_normalized)
                    coins_by_key = coins_index.get(key, [])
                    found_coin = None
                    # === ПРОХОД 1: знак МД + диаметр ===
                    for coin in coins_by_key:
                        if coin.id in processed_coin_ids:
                            continue
                        if mint_mark:
                            coin_mint = (coin.mint_mark or "").strip()
                            if coin_mint.lower() != mint_mark.lower():
                                continue
                        else:
                            if coin.mint_mark and str(coin.mint_mark).strip():
                                continue
                        if not self._diameter_matches(coin, diameter):
                            continue
                        found_coin = coin
                        break
                    # === ПРОХОД 2: диаметр + совпадение «Наименования» с описанием ===
                    if not found_coin and name_marker:
                        nm_low = name_marker.lower()
                        for coin in coins_by_key:
                            if coin.id in processed_coin_ids:
                                continue
                            if not self._diameter_matches(coin, diameter):
                                continue
                            blob = f"{getattr(coin, 'coin_info', '') or ''} " \
                                   f"{getattr(coin, 'description', '') or ''}".lower()
                            if nm_low in blob:
                                found_coin = coin
                                break
                    # === ПРОХОД 3: только диаметр ===
                    if not found_coin:
                        for coin in coins_by_key:
                            if coin.id in processed_coin_ids:
                                continue
                            if not self._diameter_matches(coin, diameter):
                                continue
                            found_coin = coin
                            break
                    # === ПРОХОД 4: страна + год + номинал ===
                    if not found_coin:
                        for coin in coins_by_key:
                            if coin.id in processed_coin_ids:
                                continue
                            found_coin = coin
                            break
                    if not found_coin:
                        error_count += 1
                        error_msg = (f"Строка {idx + 2}: монета не найдена "
                                     f"(страна={country.name}, год={year}, "
                                     f"номинал={denomination_raw})")
                        if mint_mark:
                            error_msg += f", знак МД={mint_mark}"
                        if diameter:
                            error_msg += f", диаметр={diameter}"
                        errors_list.append(error_msg)
                        self.logger.warning(f"❌ Монета не найдена для строки {idx + 2}: {error_msg}")
                        continue
                    processed_coin_ids.add(found_coin.id)
                    # === ОБНОВЛЯЕМ ТОЛЬКО ЦЕНУ И ДАТУ ЦЕНЫ ===
                    # Описание монеты (coin_info / description) НЕ изменяется.
                    found_coin.market_price = price
                    found_coin.market_price_date = self.file_date
                    self.db_manager.add_market_price_history(
                        coin_id=found_coin.id,
                        market_price=price,
                        price_date=self.file_date,
                        source_file=source_file
                    )
                    updated_count += 1
                    if updated_count % 50 == 0:
                        self.db_manager.session.commit()
                        self.progress.emit(idx + 1, total_rows,
                                           f"✅ Обработано {updated_count}...")
                except Exception as e:
                    error_count += 1
                    errors_list.append(f"Строка {idx + 2}: ошибка - {str(e)}")
                    self.logger.error(f"Ошибка в строке {idx + 2}: {e}")
            self.db_manager.session.commit()
            errors_list.sort()
            self.progress.emit(total_rows, total_rows,
                               f"✅ Успешно: {updated_count} | ❌ Ошибок: {error_count}")
            self.finished.emit(updated_count, error_count, errors_list)
        except Exception as e:
            self.logger.error(f"Ошибка: {e}")
            import traceback
            traceback.print_exc()
            self.error.emit(str(e))

    def stop(self):
        self._is_running = False


    @staticmethod
    def _diameter_matches(coin, diameter):
        """Сравнивает диаметр монеты с диаметром из прайса (с округлением до 0.1)."""
        if diameter is None:
            return True
        coin_dia = getattr(coin, 'diameter', None)
        if coin_dia is None:
            return True
        try:
            return round(float(coin_dia), 1) == round(float(diameter), 1)
        except (ValueError, TypeError):
            return str(coin_dia).strip() == str(diameter).strip()

class LoadMarketPricesDialog(QDialog):
    """Диалог для загрузки рыночных цен из Excel"""
    
    DEFAULT_PRICE_DIR = paths.get_price_dir()
    
    def __init__(self, db_manager, parent=None):
        super().__init__(parent)
        self.db_manager = db_manager
        self.load_thread = None
        self.file_date = None
        self.setWindowTitle("Загрузка рыночных цен")
        self.setMinimumSize(750, 600)
        
        self.init_ui()
        QTimer.singleShot(100, self.auto_load_latest_file)
    
    def init_ui(self):
        layout = QVBoxLayout()
        self.setLayout(layout)
        
        title = QLabel("📥 Загрузка рыночных цен из Excel")
        title.setStyleSheet("""
            font-weight: bold;
            font-size: 14px;
            padding: 8px;
            background-color: #4a6fa5;
            color: white;
            border-radius: 3px;
        """)
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)
        
        info_group = QGroupBox("📋 Формат файла Excel")
        info_layout = QGridLayout()
        
        info_layout.addWidget(QLabel("Столбец A:"), 0, 0)
        info_layout.addWidget(QLabel("Страна"), 0, 1)
        
        info_layout.addWidget(QLabel("Столбец D:"), 1, 0)
        info_layout.addWidget(QLabel("Номинал (дроби: ¼, ½, ¾)"), 1, 1)
        
        info_layout.addWidget(QLabel("Столбец E:"), 2, 0)
        info_layout.addWidget(QLabel("Год"), 2, 1)
        
        info_layout.addWidget(QLabel("Столбец F:"), 3, 0)
        info_layout.addWidget(QLabel("Знак МД"), 3, 1)
        
        info_layout.addWidget(QLabel("Столбец H:"), 4, 0)
        info_layout.addWidget(QLabel("Диаметр (мм)"), 4, 1)
        
        info_layout.addWidget(QLabel("Столбец J:"), 5, 0)
        info_layout.addWidget(QLabel("Рыночная цена"), 5, 1)
        
        info_group.setLayout(info_layout)
        layout.addWidget(info_group)
        
        file_layout = QHBoxLayout()
        file_layout.addWidget(QLabel("Файл:"))
        
        self.file_path_edit = QLabel()
        self.file_path_edit.setStyleSheet("border: 1px solid #ccc; padding: 5px; background-color: #f5f5f5;")
        self.file_path_edit.setWordWrap(True)
        file_layout.addWidget(self.file_path_edit, 1)
        
        self.browse_btn = QPushButton("📂 Обзор")
        self.browse_btn.clicked.connect(self.browse_file)
        file_layout.addWidget(self.browse_btn)
        
        layout.addLayout(file_layout)
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)
        
        log_label = QLabel("📋 Лог операций:")
        layout.addWidget(log_label)
        
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setMaximumHeight(150)
        layout.addWidget(self.log_text)
        
        errors_label = QLabel("❌ Детальные ошибки (не найдены монеты):")
        layout.addWidget(errors_label)
        
        self.errors_text = QTextEdit()
        self.errors_text.setReadOnly(True)
        self.errors_text.setMaximumHeight(200)
        self.errors_text.setPlaceholderText("Здесь будут показаны причины, по которым монеты не были найдены...")
        layout.addWidget(self.errors_text)
        
        button_layout = QHBoxLayout()
        
        self.load_btn = QPushButton("🚀 Загрузить цены")
        self.load_btn.clicked.connect(self.load_prices)
        self.load_btn.setEnabled(False)
        self.load_btn.setMinimumHeight(35)
        self.load_btn.setStyleSheet("background-color: #4CAF50; color: white; font-weight: bold;")
        button_layout.addWidget(self.load_btn)
        
        self.cancel_btn = QPushButton("❌ Отмена")
        self.cancel_btn.clicked.connect(self.cancel_loading)
        self.cancel_btn.setEnabled(False)
        self.cancel_btn.setMinimumHeight(35)
        button_layout.addWidget(self.cancel_btn)
        
        button_layout.addStretch()
        
        self.copy_errors_btn = QPushButton("📋 Копировать ошибки")
        self.copy_errors_btn.clicked.connect(self.copy_errors)
        self.copy_errors_btn.setEnabled(False)
        self.copy_errors_btn.setMinimumHeight(35)
        button_layout.addWidget(self.copy_errors_btn)
        
        self.close_btn = QPushButton("❌ Закрыть")
        self.close_btn.clicked.connect(self.accept)
        self.close_btn.setMinimumHeight(35)
        button_layout.addWidget(self.close_btn)
        
        layout.addLayout(button_layout)
    
    def find_latest_price_file(self):
        if not self.DEFAULT_PRICE_DIR.exists():
            return None
        
        price_files = list(self.DEFAULT_PRICE_DIR.glob("price_*.xls")) + list(self.DEFAULT_PRICE_DIR.glob("price_*.xlsx"))
        if price_files:
            return max(price_files, key=lambda f: f.stat().st_mtime)
        return None
    
    def extract_date_from_filename(self, filename):
        match = re.search(r'price_(\d{8})_\d{6}', filename)
        if match:
            try:
                return datetime.strptime(match.group(1), "%Y%m%d").date()
            except:
                pass
        
        match = re.search(r'(\d{4}-\d{2}-\d{2})', filename)
        if match:
            try:
                return datetime.strptime(match.group(1), "%Y-%m-%d").date()
            except:
                pass
        
        return None
    
    def auto_load_latest_file(self):
        latest_file = self.find_latest_price_file()
        if latest_file:
            self.file_path_edit.setText(str(latest_file))
            self.load_btn.setEnabled(True)
            
            self.file_date = self.extract_date_from_filename(latest_file.stem)
            
            if self.file_date:
                self.add_log(f"📁 Автоматически загружен: {latest_file.name}", "success")
                self.add_log(f"📅 Дата: {self.file_date}", "success")
            else:
                mod_time = latest_file.stat().st_mtime
                self.file_date = datetime.fromtimestamp(mod_time).date()
                self.add_log(f"📁 Автоматически загружен: {latest_file.name}", "success")
                self.add_log(f"📅 Дата модификации: {self.file_date}", "info")
            
            QTimer.singleShot(500, self.load_prices)
        else:
            self.add_log("📁 Файлы прайса не найдены. Скачайте прайс через браузер (кнопка 'Прайс')", "info")
    
    def browse_file(self):
        self.DEFAULT_PRICE_DIR.mkdir(parents=True, exist_ok=True)
        
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл с рыночными ценами",
            str(self.DEFAULT_PRICE_DIR),
            "Excel files (*.xlsx *.xls);;All files (*.*)"
        )
        
        if file_path:
            self.file_path_edit.setText(file_path)
            self.load_btn.setEnabled(True)
            
            file_name = Path(file_path).stem
            self.file_date = self.extract_date_from_filename(file_name)
            
            if self.file_date:
                self.add_log(f"📅 Дата из имени: {self.file_date}", "success")
            else:
                mod_time = Path(file_path).stat().st_mtime
                self.file_date = datetime.fromtimestamp(mod_time).date()
                self.add_log(f"📅 Дата модификации: {self.file_date}", "info")
    
    def add_log(self, message, level="info"):
        timestamp = datetime.now().strftime("%H:%M:%S")
        prefix = {"error": "❌", "success": "✅", "warning": "⚠️", "info": "ℹ️"}.get(level, "📌")
        self.log_text.append(f"[{timestamp}] {prefix} {message}")
        scrollbar = self.log_text.verticalScrollBar()
        scrollbar.setValue(scrollbar.maximum())
        QApplication.processEvents()
    
    def load_prices(self):
        file_path = self.file_path_edit.text()
        if not file_path or not Path(file_path).exists():
            QMessageBox.warning(self, "Предупреждение", "Выберите файл")
            return
        
        if not self.file_date:
            self.file_date = datetime.now().date()
            self.add_log(f"📅 Используется сегодняшняя дата: {self.file_date}", "warning")
        
        reply = QMessageBox.question(
            self, "Подтверждение",
            f"Загрузить рыночные цены?\n\n"
            f"Файл: {Path(file_path).name}\n"
            f"Дата: {self.file_date}\n\n"
            f"Продолжить?",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply != QMessageBox.Yes:
            return
        
        self.load_btn.setEnabled(False)
        self.browse_btn.setEnabled(False)
        self.cancel_btn.setEnabled(True)
        self.close_btn.setEnabled(False)
        self.copy_errors_btn.setEnabled(False)
        
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.errors_text.clear()
        
        self.load_thread = FastLoadPricesThread(self.db_manager, file_path, self.file_date)
        self.load_thread.progress.connect(self.on_progress)
        self.load_thread.finished.connect(self.on_finished)
        self.load_thread.error.connect(self.on_error)
        self.load_thread.start()
    
    def cancel_loading(self):
        if self.load_thread and self.load_thread.isRunning():
            self.load_thread.stop()
            self.add_log("Загрузка отменена пользователем", "warning")
            self.cancel_btn.setEnabled(False)
    
    def on_progress(self, current, total, message):
        percent = int((current / total) * 100) if total > 0 else 0
        self.progress_bar.setValue(percent)
        self.add_log(message, "info")
    
    def on_finished(self, updated_count, error_count, errors_list):
        self.progress_bar.setVisible(False)
        self.cancel_btn.setEnabled(False)
        self.close_btn.setEnabled(True)
        self.load_btn.setEnabled(True)
        self.browse_btn.setEnabled(True)
        
        self.add_log("=" * 50, "info")
        self.add_log(f"✅ Загрузка завершена!", "success")
        self.add_log(f"📊 Обновлено монет: {updated_count}", "success")
        
        if error_count > 0:
            self.add_log(f"⚠️ Не найдено/ошибок: {error_count}", "warning")
            self.copy_errors_btn.setEnabled(True)
            
            error_text = "\n".join(errors_list[:200])
            if len(errors_list) > 200:
                error_text += f"\n\n... и ещё {len(errors_list) - 200} ошибок"
            self.errors_text.setText(error_text)
        else:
            self.errors_text.setText("✅ Все монеты успешно найдены и обновлены!")
        
        self.add_log("=" * 50, "info")
        
        QMessageBox.information(
            self, "Загрузка завершена",
            f"✅ Обновлено монет: {updated_count}\n"
            f"❌ Не найдено/ошибок: {error_count}\n\n"
            f"Детальные ошибки показаны в окне ниже.\n"
            f"Их можно скопировать кнопкой 'Копировать ошибки'."
        )
    
    def on_error(self, error_msg):
        self.progress_bar.setVisible(False)
        self.cancel_btn.setEnabled(False)
        self.close_btn.setEnabled(True)
        self.load_btn.setEnabled(True)
        self.browse_btn.setEnabled(True)
        
        self.add_log(f"❌ Критическая ошибка: {error_msg}", "error")
        self.errors_text.setText(f"Критическая ошибка:\n{error_msg}")
        QMessageBox.critical(self, "Ошибка", f"Не удалось загрузить цены:\n{error_msg}")
    
    def copy_errors(self):
        text = self.errors_text.toPlainText()
        if text:
            from PySide6.QtWidgets import QApplication
            clipboard = QApplication.clipboard()
            clipboard.setText(text)
            self.add_log("📋 Ошибки скопированы в буфер обмена", "success")
            
            