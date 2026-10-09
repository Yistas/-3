# -*- coding: utf-8 -*-

"""Загрузка и сохранение Excel файлов"""

import os
import logging
import pandas as pd
from datetime import datetime
from PySide6.QtWidgets import QMessageBox, QApplication, QFileDialog
from PySide6.QtCore import QTimer
from PySide6.QtGui import QBrush, QColor

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelIO:
    """Класс для работы с Excel файлами"""
    
    def __init__(self):
        self.current_workbook = None
        self.logger = logging.getLogger('CoinCollector.GUI.ExcelIO')
    
    def load_excel_with_formatting(self, file_path):
        """Загружает Excel файл с форматированием"""
        if not OPENPYXL_AVAILABLE:
            self._load_without_formatting(file_path)
            return
        
        try:
            wb = load_workbook(file_path, data_only=False)
            self.current_workbook = wb
            
            self.sheet_tabs.clear()
            self.sheets.clear()
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                self.add_new_sheet(sheet_name)
                sheet = self.sheets[sheet_name]
                table = sheet.table
                table.sheet_name = sheet_name
                
                max_row = ws.max_row
                max_col = ws.max_column
                
                self.logger.info(f"Загружается лист '{sheet_name}': строк={max_row}, колонок={max_col}")
                
                if max_row > 0 and max_col > 0:
                    table.setRowCount(max_row - 1)
                    table.setColumnCount(max_col)
                    
                    # Заголовки
                    for col in range(1, max_col + 1):
                        header_cell = ws.cell(row=1, column=col)
                        header_text = str(header_cell.value) if header_cell.value else f"Колонка {col}"
                        header_item = QTableWidgetItem(header_text)
                        self._apply_cell_formatting(header_item, header_cell)
                        table.setHorizontalHeaderItem(col - 1, header_item)
                    
                    # Данные
                    for row in range(2, max_row + 1):
                        for col in range(1, max_col + 1):
                            try:
                                cell = ws.cell(row=row, column=col)
                                value = cell.value
                                item = QTableWidgetItem()
                                
                                if isinstance(value, str) and value.startswith('='):
                                    item.setData(Qt.UserRole + 1, value)
                                    result = table.evaluate_formula(value, row-2, col-1)
                                    item.setText(str(result) if result else "#ERROR")
                                    item.setForeground(QBrush(QColor(0, 100, 0)))
                                elif value is not None:
                                    if isinstance(value, (int, float)):
                                        item.setText(f"{value:.2f}" if isinstance(value, float) else str(value))
                                    elif isinstance(value, datetime):
                                        item.setText(value.strftime("%Y-%m-%d"))
                                    else:
                                        item.setText(str(value))
                                else:
                                    item.setText("")
                                
                                self._apply_cell_formatting(item, cell)
                                table.setItem(row - 2, col - 1, item)
                            except Exception as e:
                                self.logger.error(f"Ошибка ячейки [{row},{col}]: {e}")
                        
                        if row % 100 == 0:
                            QApplication.processEvents()
                    
                    # Восстанавливаем размеры
                    table.restore_table_dimensions(sheet_name)
                    
                    # Ширина из файла
                    for col in range(1, max_col + 1):
                        try:
                            col_letter = get_column_letter(col)
                            col_width = ws.column_dimensions[col_letter].width
                            if col_width and col_width > 0:
                                table.setColumnWidth(col - 1, int(col_width * 7))
                        except:
                            pass
                    
                    self.logger.info(f"✅ Загружен лист '{sheet_name}': {max_row-1} строк, {max_col} колонок")
            
            self.is_modified = False
            
        except Exception as e:
            self.logger.error(f"Ошибка загрузки: {e}")
            self._load_without_formatting(file_path)
    
    def save_with_formatting(self, file_path):
        """Сохраняет таблицу в Excel с форматированием"""
        if not OPENPYXL_AVAILABLE:
            self.save_to_file(file_path)
            return
        
        try:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
            
            for sheet_name, sheet in self.sheets.items():
                ws = wb.create_sheet(sheet_name)
                table = sheet.table
                table.sheet_name = sheet_name
                
                # Заголовки
                for col in range(table.columnCount()):
                    header = table.horizontalHeaderItem(col)
                    header_text = header.text() if header else f"Колонка {col+1}"
                    cell = ws.cell(row=1, column=col+1, value=header_text)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="4a6fa5", end_color="4a6fa5", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                
                # Данные
                for row in range(table.rowCount()):
                    for col in range(table.columnCount()):
                        item = table.item(row, col)
                        cell = ws.cell(row=row+2, column=col+1)
                        
                        if item:
                            formula = item.data(Qt.UserRole + 1)
                            if formula:
                                cell.value = formula
                            else:
                                cell.value = item.text()
                            
                            self._apply_item_formatting_to_cell(item, cell)
                
                # Ширина колонок
                for col in range(table.columnCount()):
                    col_width = table.columnWidth(col)
                    if col_width > 0:
                        col_letter = get_column_letter(col + 1)
                        ws.column_dimensions[col_letter].width = col_width / 7
                
                table.save_column_widths(sheet_name)
            
            wb.save(file_path)
            self.current_workbook = wb
            self.logger.info(f"Файл сохранен: {file_path}")
            
        except Exception as e:
            self.logger.error(f"Ошибка сохранения: {e}")
            raise
    
    def _load_without_formatting(self, file_path):
        """Загружает без форматирования"""
        try:
            df = pd.read_excel(file_path)
            self.load_dataframe(df)
            QMessageBox.information(self, "Информация", "Файл загружен без форматирования")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось загрузить файл:\n{e}")
    
    def load_dataframe(self, df):
        """Загружает данные из DataFrame"""
        self.sheet_tabs.clear()
        self.sheets.clear()
        self.add_new_sheet("Лист1")
        
        sheet = self.sheets["Лист1"]
        table = sheet.table
        
        table.setRowCount(len(df))
        table.setColumnCount(len(df.columns))
        
        for i, col in enumerate(df.columns):
            table.setHorizontalHeaderItem(i, QTableWidgetItem(str(col)))
        
        for i, row in df.iterrows():
            for j, val in enumerate(row):
                table.setItem(i, j, QTableWidgetItem(str(val) if pd.notna(val) else ""))
    
    def save_to_file(self, file_path):
        """Сохраняет в файл без форматирования"""
        try:
            table = self.get_current_table()
            if not table:
                return
            
            data, headers = [], []
            for col in range(table.columnCount()):
                header = table.horizontalHeaderItem(col)
                headers.append(header.text() if header else f"Col{col+1}")
            
            for row in range(table.rowCount()):
                row_data = []
                for col in range(table.columnCount()):
                    item = table.item(row, col)
                    row_data.append(item.text() if item else "")
                data.append(row_data)
            
            df = pd.DataFrame(data, columns=headers)
            
            if file_path.endswith('.csv'):
                df.to_csv(file_path, index=False, encoding='utf-8-sig')
            else:
                with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                    df.to_excel(writer, sheet_name='Лист1', index=False)
            
            self.is_modified = False
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить:\n{e}")
    
    def _apply_cell_formatting(self, item, cell):
        """Применяет форматирование из Excel к QTableWidgetItem"""
        try:
            if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
                rgb = cell.fill.fgColor.rgb
                if rgb and rgb != '00000000':
                    color_hex = rgb[2:] if len(rgb) == 8 else rgb
                    color = QColor(f"#{color_hex}")
                    if color.isValid():
                        item.setBackground(QBrush(color))
            
            if cell.font:
                font = QFont()
                if cell.font.name:
                    font.setFamily(cell.font.name)
                if cell.font.sz:
                    font.setPointSize(int(cell.font.sz))
                if cell.font.b:
                    font.setBold(True)
                if cell.font.i:
                    font.setItalic(True)
                if cell.font.u:
                    font.setUnderline(True)
                item.setFont(font)
            
            if cell.alignment:
                if cell.alignment.horizontal == 'center':
                    item.setTextAlignment(Qt.AlignCenter)
                elif cell.alignment.horizontal == 'right':
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                else:
                    item.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
                
                if cell.alignment.vertical == 'center':
                    current = item.textAlignment()
                    horizontal = current & Qt.AlignHorizontal_Mask
                    item.setTextAlignment(horizontal | Qt.AlignVCenter)
                elif cell.alignment.vertical == 'bottom':
                    current = item.textAlignment()
                    horizontal = current & Qt.AlignHorizontal_Mask
                    item.setTextAlignment(horizontal | Qt.AlignBottom)
        except Exception:
            pass
    
    def _apply_item_formatting_to_cell(self, item, cell):
        """Применяет форматирование из QTableWidgetItem к Excel ячейке"""
        from openpyxl.styles import PatternFill, Font, Alignment
        
        if item.background().color().isValid():
            color = item.background().color().name().lstrip('#')
            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        
        font_props = {}
        if item.font().bold():
            font_props['bold'] = True
        if item.font().italic():
            font_props['italic'] = True
        if item.font().underline():
            font_props['underline'] = 'single'
        if item.font().family():
            font_props['name'] = item.font().family()
        
        if item.foreground().color().isValid():
            font_pro