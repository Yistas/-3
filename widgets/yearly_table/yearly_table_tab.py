# ===== gui/widgets/yearly_table/yearly_table_tab.py =====
# ПОЛНОСТЬЮ ЗАМЕНИТЬ ФАЙЛ

# -*- coding: utf-8 -*-

"""
Вкладка для работы с погодовкой (таблица по годам)
"""

import os
import logging
import re
import json
from datetime import datetime
from pathlib import Path

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView,
    QFileDialog, QMessageBox, QComboBox, QLabel,
    QFrame, QSplitter, QTextEdit, QColorDialog,
    QFontDialog, QMenu, QInputDialog, QLineEdit,
    QToolBar, QSizePolicy, QApplication, QTabWidget,
    QProgressBar, QProgressDialog
)
from PySide6.QtCore import Qt, Signal, QTimer, QMimeData, QSettings, QThread
from PySide6.QtGui import QAction, QColor, QFont, QBrush, QKeySequence, QPainter, QPen

from utils.paths import paths

# Для работы с Excel
try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .sheet_widget import SheetWidget
from .excel_like_table import ExcelLikeTable
from .color_settings import ColorSettings
from .color_settings_dialog import ColorSettingsDialog


class LoadTableThread(QThread):
    """Поток для фоновой загрузки таблицы"""
    finished = Signal(object)
    error = Signal(str)
    
    def __init__(self, db_manager, country_id, country_name):
        super().__init__()
        self.db_manager = db_manager
        self.country_id = country_id
        self.country_name = country_name
        self._is_running = True


        
    def run(self):
        try:
            if not self._is_running:
                return
            
            yearly_dir = paths.get_yearly_tables_dir()
            
            if self.country_id is not None:
                file_path = yearly_dir / f"country_{self.country_id}.xlsx"
            else:
                if self.country_name:
                    safe_name = "".join(c for c in self.country_name if c.isalnum() or c in ' _-')
                    file_path = yearly_dir / f"{safe_name}.xlsx"
                else:
                    file_path = yearly_dir / "all_countries.xlsx"
            
            if not self._is_running:
                return
            
            result = {
                'file_path': file_path,
                'exists': file_path.exists(),
                'country_id': self.country_id,
                'country_name': self.country_name
            }
            self.finished.emit(result)
        except Exception as e:
            if self._is_running:
                self.error.emit(str(e))
    
    def stop(self):
        self._is_running = False


class YearlyTableTab(QWidget):
    """Вкладка для работы с погодовкой (таблица по годам)"""
    
    def __init__(self, db_manager, parent=None):
        super().__init__(parent)
        self.db_manager = db_manager
        self.logger = logging.getLogger('CoinCollector.GUI.YearlyTable')
        
        self.current_file = None
        self.current_country_id = None
        self.current_country_name = None
        self._loaded_country_id = None
        self._loaded_country_name = None
        self._loading_table = False
        self.current_data = None
        self.is_modified = False
        self.current_workbook = None
        self.sheets = {}
        self._loading_country_id = None
        self._loading_country_name = None
        
        # Кэш загруженных погодовок
        self.loaded_tables_cache = {}
        self.current_load_thread = None
        self.is_loading = False
        self.pending_country_id = None
        self.pending_country_name = None
        
        # Настройки автосохранения
        self.auto_save_enabled = True
        self.auto_save_interval = 5 * 60 * 1000
        
        # Настройки цветов
        self.color_settings = ColorSettings()
        
        # Папка для сохранения файлов
        self.yearly_dir = paths.get_yearly_tables_dir()
        self.yearly_dir.mkdir(parents=True, exist_ok=True)
        
        # Настройки для сохранения порядка вкладок
        self.settings = QSettings('CoinCollector', 'YearlyTable')
        
        self.init_ui()
        self.create_empty_table()
        
        # Таймер для автосохранения
        self.auto_save_timer = QTimer()
        self.auto_save_timer.timeout.connect(self.auto_save)
        self.auto_save_timer.start(self.auto_save_interval)

    def _get_theme(self):
        """Возвращает словарь текущей темы (поднимаясь к главному окну)"""
        w = self.parent()
        while w is not None:
            tm = getattr(w, 'theme_manager', None)
            if tm is not None:
                return tm.current_theme
            w = w.parent()
        return {}
    
    def init_ui(self):
        """Инициализация интерфейса (тёмная тема)"""
        from PySide6.QtCore import QTimer
        t = self._get_theme()
        layout = QVBoxLayout()
        layout.setContentsMargins(5, 5, 5, 5)
        layout.setSpacing(5)
        self.setLayout(layout)
        # Перехватываем SAVE ALL: сначала синхронизируем порядок листов, потом пишем кэш
        self._orig_save_all_sheets = self.save_all_sheets
        self.save_all_sheets = self._save_all_with_order
        
        # Заголовок — self.title_label (используется в _update_title)
        self.title_label = QLabel("📅 Погодовка")
        self.title_label.setStyleSheet(f"""
            font-weight: bold;
            font-size: 14px;
            padding: 6px;
            background-color: {t.get('surface', '#16213e')};
            color: {t.get('accent', '#6c63ff')};
            border: 1px solid {t.get('border', '#2a2a4a')};
            border-radius: 6px;
        """)
        self.title_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(self.title_label)

        # Панель инструментов
        toolbar = self.create_toolbar()
        layout.addWidget(toolbar)

        # ВАЖНО: вкладки листов создаём ДО легенды —
        # create_color_legend() обращается к self.sheet_tabs
        self.sheet_tabs = QTabWidget()
        self.sheet_tabs.setTabsClosable(True)
        self.sheet_tabs.setMovable(True)
        self.sheet_tabs.tabCloseRequested.connect(self.close_sheet)
        self.sheet_tabs.currentChanged.connect(self.on_tab_changed)
                # Контекстное меню вкладок (переименование / закрытие)
        self.sheet_tabs.tabBar().setContextMenuPolicy(Qt.CustomContextMenu)
        self.sheet_tabs.tabBar().customContextMenuRequested.connect(self._show_sheet_context_menu)
        # Следим за ручным перетаскиванием вкладок
        self.sheet_tabs.tabBar().tabMoved.connect(self._on_sheet_tab_moved)
        self.tab_widget = self.sheet_tabs  # алиас для совместимости

        # Легенда цветов (теперь sheet_tabs уже существует)
        self.legend_widget = QFrame()
        self.legend_widget.setFrameShape(QFrame.StyledPanel)
        self.legend_widget.setStyleSheet(f"""
            QFrame {{
                background-color: {t.get('base', '#1a1a2e')};
                border: 1px solid {t.get('border', '#2a2a4a')};
                border-radius: 6px;
                padding: 2px;
            }}
        """)
        self.legend_layout = QHBoxLayout()
        self.legend_layout.setContentsMargins(8, 3, 8, 3)
        self.legend_layout.setSpacing(10)
        self.legend_widget.setLayout(self.legend_layout)
        try:
            self.create_color_legend()
            self._style_legend_dark()
        except Exception as e:
            self.logger.error(f"Ошибка создания легенды цветов: {e}")

        # Порядок в layout: легенда ВЫШЕ вкладок (как было изначально)
        layout.addWidget(self.legend_widget)
        layout.addWidget(self.sheet_tabs)

        # Нижняя информационная строка
        bottom = QWidget()
        bottom_layout = QHBoxLayout()
        bottom_layout.setContentsMargins(0, 2, 0, 2)
        bottom_layout.setSpacing(8)
        bottom.setLayout(bottom_layout)
        self.status_label = QLabel("Готов")
        self.status_label.setStyleSheet(f"color: {t.get('tab_text', '#8888aa')};")
        bottom_layout.addWidget(self.status_label)
        bottom_layout.addStretch()
        self.cell_info_label = QLabel("")
        self.cell_info_label.setStyleSheet(f"color: {t.get('tab_text', '#8888aa')};")
        bottom_layout.addWidget(self.cell_info_label)
        layout.addWidget(bottom)

        # Таймер отложенного обновления легенды
        self.legend_update_timer = QTimer()
        self.legend_update_timer.setSingleShot(True)
        self.legend_update_timer.timeout.connect(self._rebuild_color_legend)

    def showEvent(self, event):
        """Загружаем погодовку когда вкладка становится видимой"""
        super().showEvent(event)
        if self.is_loading or self._loading_table:
            return
        if self.current_country_id is None:
            return
        if self.current_country_id == self._loaded_country_id:
            return
        if self._loaded_country_id is None:
            self._load_country_table()
        elif self.current_country_id != self._loaded_country_id:
            self._load_country_table()
        QTimer.singleShot(0, self._restore_sheet_order)

    def create_toolbar(self):
        """Создает панель инструментов (тёмная тема, без кнопки высоты строк)"""
        t = self._get_theme()
        toolbar = QWidget()
        toolbar.setStyleSheet(f"""
            QWidget {{
                background-color: {t.get('base', '#1a1a2e')};
                border: 1px solid {t.get('border', '#2a2a4a')};
                border-radius: 6px;
                padding: 2px;
            }}
        """)
        layout = QHBoxLayout()
        layout.setContentsMargins(5, 3, 5, 3)
        layout.setSpacing(3)
        btn_style = f"""
            QPushButton {{
                background-color: {t.get('button', '#16213e')};
                color: {t.get('text', '#e4e4ef')};
                border: 1px solid {t.get('border', '#2a2a4a')};
                border-radius: 6px;
            }}
            QPushButton:hover {{
                background-color: {t.get('accent', '#6c63ff')};
                color: #ffffff;
                border-color: {t.get('accent', '#6c63ff')};
            }}
        """
        new_btn = QPushButton("📄")
        new_btn.clicked.connect(self.create_new_table)
        new_btn.setFixedSize(28, 28)
        new_btn.setToolTip("Создать новую таблицу")
        new_btn.setStyleSheet(btn_style)
        layout.addWidget(new_btn)
        open_btn = QPushButton("📂")
        open_btn.clicked.connect(self.open_file)
        open_btn.setFixedSize(28, 28)
        open_btn.setToolTip("Открыть файл Excel")
        open_btn.setStyleSheet(btn_style)
        layout.addWidget(open_btn)
        save_all_btn = QPushButton("💾 SAVE ALL")
        save_all_btn.clicked.connect(self.save_all_sheets)
        save_all_btn.clicked.connect(self._save_sheet_order)
        save_all_btn.setMinimumHeight(28)
        save_all_btn.setMinimumWidth(80)
        save_all_btn.setToolTip("Сохранить все изменения")
        save_all_btn.setStyleSheet(btn_style)
        layout.addWidget(save_all_btn)
        auto_fit_btn = QPushButton("↔️ Автоширина")
        auto_fit_btn.clicked.connect(self.auto_fit_columns)
        auto_fit_btn.setMinimumHeight(28)
        auto_fit_btn.setMinimumWidth(90)
        auto_fit_btn.setToolTip("Подобрать ширину колонок по содержимому и центрировать текст")
        auto_fit_btn.setStyleSheet("""
            QPushButton { background-color: #17a2b8; color: white; font-weight: bold; border-radius: 6px; }
            QPushButton:hover { background-color: #138496; }
        """)
        layout.addWidget(auto_fit_btn)
        borders_btn = QPushButton("🔲")
        borders_btn.clicked.connect(self.show_borders_menu)
        borders_btn.setMinimumHeight(28)
        borders_btn.setMinimumWidth(80)
        borders_btn.setToolTip("Границы ячеек")
        borders_btn.setStyleSheet("""
            QPushButton { background-color: #6c757d; color: white; font-weight: bold; border-radius: 6px; }
            QPushButton:hover { background-color: #5a6268; }
        """)
        layout.addWidget(borders_btn)
        save_as_btn = QPushButton("📎")
        save_as_btn.clicked.connect(self.save_file_as)
        save_as_btn.setFixedSize(28, 28)
        save_as_btn.setToolTip("Сохранить в новый файл")
        save_as_btn.setStyleSheet(btn_style)
        layout.addWidget(save_as_btn)
        separator1 = QFrame()
        separator1.setFrameShape(QFrame.VLine)
        separator1.setFrameShadow(QFrame.Sunken)
        separator1.setFixedSize(2, 24)
        layout.addWidget(separator1)
        merge_btn = QPushButton("🔗")
        merge_btn.clicked.connect(self.merge_selected_cells)
        merge_btn.setFixedSize(28, 28)
        merge_btn.setToolTip("Объединить выделенные ячейки")
        merge_btn.setStyleSheet(btn_style)
        layout.addWidget(merge_btn)
        unmerge_btn = QPushButton("🔓")
        unmerge_btn.clicked.connect(self.unmerge_selected_cells)
        unmerge_btn.setFixedSize(28, 28)
        unmerge_btn.setToolTip("Разъединить выделенные ячейки")
        unmerge_btn.setStyleSheet(btn_style)
        layout.addWidget(unmerge_btn)
        separator2 = QFrame()
        separator2.setFrameShape(QFrame.VLine)
        separator2.setFrameShadow(QFrame.Sunken)
        separator2.setFixedSize(2, 24)
        layout.addWidget(separator2)
        import_ucoin_btn = QPushButton("🌐 Импорт из UCOIN")
        import_ucoin_btn.clicked.connect(self.import_from_ucoin)
        import_ucoin_btn.setMinimumHeight(28)
        import_ucoin_btn.setMinimumWidth(120)
        import_ucoin_btn.setToolTip("Импортировать таблицу монет из открытой вкладки браузера uCoin.net")
        import_ucoin_btn.setStyleSheet("""
            QPushButton { background-color: #e67e22; color: white; font-weight: bold; border-radius: 6px; }
            QPushButton:hover { background-color: #f39c12; }
        """)
        layout.addWidget(import_ucoin_btn)
        separator3 = QFrame()
        separator3.setFrameShape(QFrame.VLine)
        separator3.setFrameShadow(QFrame.Sunken)
        separator3.setFixedSize(2, 24)
        layout.addWidget(separator3)
        color_widget = self.create_color_buttons()
        layout.addWidget(color_widget)
        layout.addStretch()
        # Индикатор загрузки
        self.loading_label = QLabel("")
        self.loading_label.setStyleSheet(f"color: {t.get('accent', '#6c63ff')}; font-size: 10px;")
        layout.addWidget(self.loading_label)
        # Метка размера СОЗДАЁТСЯ, но НЕ выводится на панель
        # (атрибут нужен другим методам для обновления текста)
        self.size_label = QLabel("")
        self.size_label.hide()
        toolbar.setLayout(layout)
        return toolbar

    def create_color_buttons(self):
        """Создает кнопки для выбора цвета заливки (тёмная тема)"""
        t = self._get_theme()
        border_light = t.get('scrollbar_handle', '#3a3a5a')
        color_widget = QWidget()
        color_layout = QHBoxLayout()
        color_layout.setContentsMargins(0, 0, 0, 0)
        color_layout.setSpacing(2)

        settings_btn = QPushButton("⚙️")
        settings_btn.setFixedSize(28, 28)
        settings_btn.setToolTip("Настройка цветов")
        settings_btn.clicked.connect(self.open_color_settings)
        settings_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {t.get('button', '#16213e')};
                color: {t.get('text', '#e4e4ef')};
                border: 1px solid {border_light};
                border-radius: 6px;
            }}
            QPushButton:hover {{ border-color: {t.get('accent', '#6c63ff')}; }}
        """)
        color_layout.addWidget(settings_btn)

        sep = QFrame()
        sep.setFrameShape(QFrame.VLine)
        sep.setFrameShadow(QFrame.Sunken)
        sep.setFixedSize(2, 24)
        color_layout.addWidget(sep)

        label = QLabel("🎨")
        label.setStyleSheet("font-size: 14px; padding: 0 2px;")
        label.setToolTip("Заливка ячеек")
        color_layout.addWidget(label)

        enabled_colors = self.color_settings.get_enabled_colors()
        for color_name, color_data in enabled_colors.items():
            color_code = color_data["code"]
            description = color_data.get("description", color_name)
            btn = QPushButton()
            btn.setFixedSize(24, 24)
            btn.setToolTip(f"{color_name}: {description}")
            btn.setCursor(Qt.PointingHandCursor)
            btn.setStyleSheet(f"""
                QPushButton {{
                    background-color: {color_code};
                    border: 1px solid {border_light};
                    border-radius: 3px;
                }}
                QPushButton:hover {{
                    border: 2px solid #ffffff;
                }}
            """)
            btn.clicked.connect(lambda checked, c=color_code: self.set_selected_cells_color(c))
            color_layout.addWidget(btn)

        clear_btn = QPushButton("✖")
        clear_btn.setFixedSize(24, 24)
        clear_btn.setToolTip("Убрать заливку")
        clear_btn.setCursor(Qt.PointingHandCursor)
        clear_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {t.get('button', '#16213e')};
                border: 1px solid {border_light};
                border-radius: 3px;
                font-size: 12px;
                color: {t.get('text', '#e4e4ef')};
            }}
            QPushButton:hover {{
                border: 2px solid #ffffff;
                background-color: {t.get('surface_hover', '#1f2b47')};
            }}
        """)
        clear_btn.clicked.connect(lambda: self.set_selected_cells_color(None))
        color_layout.addWidget(clear_btn)

        color_layout.addStretch()
        color_widget.setLayout(color_layout)
        return color_widget

    def create_color_legend(self):
        """Создает панель легенды цветов с процентами от всех цветных ячеек на листе"""
        if self.legend_widget is not None:
            try:
                self.layout().removeWidget(self.legend_widget)
                self.legend_widget.deleteLater()
            except:
                pass
        
        legend_widget = QFrame()
        legend_widget.setFrameShape(QFrame.StyledPanel)
        legend_widget.setStyleSheet("""
            QFrame {
                background-color: #f9f9f9;
                border: 1px solid #ddd;
                border-radius: 3px;
                padding: 2px;
            }
        """)
        
        legend_layout = QHBoxLayout()
        legend_layout.setContentsMargins(5, 2, 5, 2)
        legend_layout.setSpacing(8)
        
        legend_layout.addWidget(QLabel("📊 Легенда:"))
        
        table = self.get_current_table()
        
        all_color_counts = {}
        total_colored_cells = 0
        if table:
            all_color_counts = self._count_colored_cells(table)
            for color_name, color_data in self.color_settings.colors.items():
                color_code = color_data["code"]
                count_percentage = color_data.get("count_percentage", True)
                if count_percentage and color_code in all_color_counts:
                    total_colored_cells += all_color_counts[color_code]
        
        legend_colors = self.color_settings.get_legend_colors()
        
        for color_name, color_data in legend_colors.items():
            color_code = color_data["code"]
            description = color_data.get("description", color_name)
            
            color_box = QLabel()
            color_box.setFixedSize(14, 14)
            color_box.setStyleSheet(f"""
                background-color: {color_code};
                border: 1px solid #999;
                border-radius: 2px;
            """)
            legend_layout.addWidget(color_box)
            
            count = all_color_counts.get(color_code, 0)
            count_percentage = self.color_settings.should_count_percentage(color_name)
            
            if count_percentage and total_colored_cells > 0 and count > 0:
                percent = (count / total_colored_cells) * 100
                label_text = f"{description}: {count} ({percent:.1f}%)"
            elif count > 0:
                label_text = f"{description}: {count}"
            else:
                continue
            
            label = QLabel(label_text)
            label.setStyleSheet("font-size: 10px; font-weight: bold;")
            label.setToolTip(f"{color_name}: {description}")
            legend_layout.addWidget(label)
        
        legend_layout.addStretch()
        legend_widget.setLayout(legend_layout)
        
        try:
            self.layout().insertWidget(2, legend_widget)
        except:
            self.layout().addWidget(legend_widget)
        
        self.legend_widget = legend_widget
 
    def _style_legend_dark(self):
        """Перекрашивает легенду цветов под тёмную тему"""
        t = self._get_theme()
        w = getattr(self, 'legend_widget', None)
        if w is None:
            return
        w.setStyleSheet(f"""
            background-color: {t.get('base', '#1a1a2e')};
            border: 1px solid {t.get('border', '#2a2a4a')};
            border-radius: 6px;
            padding: 2px;
        """)
        # Серые подписи делаем читаемыми на тёмном фоне
        from PySide6.QtWidgets import QLabel
        for lbl in w.findChildren(QLabel):
            ss = lbl.styleSheet()
            if ss and ('#666' in ss or '#999' in ss or 'gray' in ss):
                lbl.setStyleSheet(
                    ss.replace('#666', '#a8a8d0').replace('#999', '#a8a8d0').replace('gray', '#a8a8d0')
                )
 
    def _rebuild_color_legend(self):
        """Перенаполняет легенду цветов и красит её под тёмную тему"""
        try:
            self.create_color_legend()
            self._style_legend_dark()
        except Exception as e:
            self.logger.error(f"Ошибка обновления легенды цветов: {e}")

    def _delayed_legend_update(self):
        """Отложенное обновление легенды"""
        self.create_color_legend()
    
    def refresh_color_legend(self):
        """Обновляет легенду цветов с пересчетом процентов (отложенно)"""
        self.legend_update_timer.start(100)
    
    def refresh_color_buttons(self):
        """Обновляет цветовые кнопки на панели после изменения настроек"""
        from PySide6.QtCore import QTimer
        
        def do_update():
            for i in range(self.toolbar_layout.count()):
                item = self.toolbar_layout.itemAt(i)
                if item is None:
                    continue
                widget = item.widget()
                if widget and hasattr(widget, 'layout') and widget.layout():
                    for j in range(widget.layout().count()):
                        child_item = widget.layout().itemAt(j)
                        if child_item and child_item.widget():
                            child = child_item.widget()
                            if child and child.toolTip() == "Настройка цветов":
                                self.toolbar_layout.removeWidget(widget)
                                widget.deleteLater()
                                break
                    break
            
            new_color_widget = self.create_color_buttons()
            
            insert_pos = -1
            for i in range(self.toolbar_layout.count()):
                item = self.toolbar_layout.itemAt(i)
                if item is None:
                    continue
                widget = item.widget()
                if widget and isinstance(widget, QPushButton) and widget.text() == "🔓":
                    insert_pos = i + 1
                    break
            
            if insert_pos >= 0:
                self.toolbar_layout.insertWidget(insert_pos, new_color_widget)
            else:
                self.toolbar_layout.insertWidget(self.toolbar_layout.count() - 1, new_color_widget)
            
            self.toolbar_layout.update()
        
        QTimer.singleShot(100, do_update)
    
    def open_color_settings(self):
        """Открывает диалог настройки цветов"""
        dialog = ColorSettingsDialog(self.color_settings, self)
        dialog.colors_changed.connect(self.refresh_color_buttons)
        dialog.colors_changed.connect(self.refresh_color_legend)
        
        if dialog.exec():
            pass
        else:
            self.color_settings.load_colors()
            self.refresh_color_buttons()
            self.refresh_color_legend()
    
    def set_selected_cells_color(self, color_code):
        """Устанавливает цвет заливки для выделенных ячеек"""
        table = self.get_current_table()
        if not table:
            QMessageBox.warning(self, "Предупреждение", "Нет активной таблицы")
            return
        
        if color_code is None:
            color = QColor(255, 255, 255)
        else:
            color = QColor(color_code)
        
        table.set_selected_cells_color(color)
        self._color_counts_cache = {}
        self.refresh_color_legend()
    
    def merge_selected_cells(self):
        """Объединяет выделенные ячейки в текущей таблице"""
        table = self.get_current_table()
        if table:
            table.merge_selected_cells()
    
    def unmerge_selected_cells(self):
        """Разъединяет выделенные ячейки в текущей таблице"""
        table = self.get_current_table()
        if table:
            table.unmerge_selected_cells()
    
    def set_all_rows_height_dialog(self):
        """Открывает диалог для установки высоты всех строк"""
        table = self.get_current_table()
        if not table:
            QMessageBox.warning(self, "Предупреждение", "Нет активной таблицы")
            return
        
        current_height = 20
        if table.rowCount() > 0:
            current_height = table.rowHeight(0)
        
        height, ok = QInputDialog.getInt(
            self,
            "Высота строк",
            "Введите высоту строк в пикселях (5-200):",
            current_height,
            5,
            200,
            1
        )
        
        if ok and height > 0:
            table.set_all_rows_height(height)
            sheet_name = self.sheet_tabs.tabText(self.sheet_tabs.currentIndex())
            table.save_table_dimensions(sheet_name)
            self.status_label.setText(f"✅ Высота всех строк установлена: {height}px")

    def get_current_table(self):
        """Возвращает текущую таблицу"""
        current_widget = self.sheet_tabs.currentWidget()
        if current_widget:
            table = current_widget.table
            return table
        return None

    def _count_colored_cells(self, table):
        """Подсчет цветных ячеек"""
        if not table:
            return {}
        
        color_counts = {}
        rows = table.rowCount()
        cols = table.columnCount()
        
        try:
            for row in range(rows):
                for col in range(cols):
                    item = table.item(row, col)
                    if item and item.background().color().isValid():
                        color = item.background().color()
                        color_code = color.name()
                        
                        if color_code == "#ffffff":
                            continue
                        
                        if color_code not in color_counts:
                            color_counts[color_code] = 0
                        color_counts[color_code] += 1
        except Exception as e:
            self.logger.error(f"Ошибка при подсчете цветных ячеек: {e}")
        
        return color_counts

    def add_new_sheet(self, sheet_name=None):
        """Добавляет новый лист"""
        if not sheet_name:
            sheet_name, ok = QInputDialog.getText(
                self, "Новый лист",
                "Введите название листа:",
                text=f"Лист{self.sheet_tabs.count() + 1}"
            )
            if not ok or not sheet_name:
                return
        
        sheet = SheetWidget(sheet_name)
        table = sheet.table
        table.sheet_name = sheet_name
        table.setAlternatingRowColors(False)
        
        table.horizontalHeader().sectionResized.connect(self.on_column_width_changed)
        table.verticalHeader().sectionResized.connect(self.on_row_height_changed)
        
        self.sheets[sheet_name] = sheet
        self.sheet_tabs.addTab(sheet, sheet_name)
        self.sheet_tabs.setCurrentWidget(sheet)
        
        table.dataChanged.connect(self.on_data_changed)
        table.currentCellChanged.connect(self.update_cell_info)
        
        table.restore_table_dimensions(sheet_name)
        
        self.is_modified = True
        self.status_label.setText(f"🔴 Добавлен лист: {sheet_name}")
        self.refresh_color_legend()
        self.save_tab_order()

    def close_sheet(self, index):
        """Закрывает лист"""
        if self.sheet_tabs.count() <= 1:
            QMessageBox.warning(self, "Предупреждение", "Нельзя закрыть последний лист")
            return
        
        widget = self.sheet_tabs.widget(index)
        sheet_name = self.sheet_tabs.tabText(index)
        
        if hasattr(widget, 'table'):
            widget.table.save_table_dimensions(sheet_name)
        
        reply = QMessageBox.question(
            self, "Подтверждение",
            f"Закрыть лист '{sheet_name}'?",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            if sheet_name in self.sheets:
                del self.sheets[sheet_name]
            self.sheet_tabs.removeTab(index)
            widget.deleteLater()
            self.is_modified = True
            self.status_label.setText(f"🔴 Закрыт лист: {sheet_name}")
            self.refresh_color_legend()
            self.save_tab_order()

    def on_tab_changed(self, index):
        """Обработчик смены вкладки - сохраняем размеры всех листов и обновляем легенду"""
        if index < 0:
            return
        
        self.save_tab_order()
        
        for sheet_name, sheet in self.sheets.items():
            if hasattr(sheet, 'table'):
                try:
                    sheet.table.save_table_dimensions(sheet_name)
                except Exception as e:
                    self.logger.debug(f"Ошибка сохранения размеров для {sheet_name}: {e}")
        
        self.update_size_info()
        self.refresh_color_legend()

    def on_column_width_changed(self, logical_index, old_width, new_width):
        """Обработчик изменения ширины колонки"""
        current = self.sheet_tabs.currentWidget()
        if not current or not hasattr(current, 'table'):
            return
        
        table = current.table
        sheet_name = getattr(table, 'sheet_name', None)
        
        if not sheet_name:
            idx = self.sheet_tabs.currentIndex()
            if idx >= 0:
                sheet_name = self.sheet_tabs.tabText(idx)
                table.sheet_name = sheet_name
        
        if sheet_name:
            table.save_table_dimensions(sheet_name)
    
    def on_row_height_changed(self, logical_index, old_height, new_height):
        """Обработчик изменения высоты строки"""
        current = self.sheet_tabs.currentWidget()
        if not current or not hasattr(current, 'table'):
            return
        
        table = current.table
        sheet_name = getattr(table, 'sheet_name', None)
        
        if not sheet_name:
            idx = self.sheet_tabs.currentIndex()
            if idx >= 0:
                sheet_name = self.sheet_tabs.tabText(idx)
                table.sheet_name = sheet_name
        
        if sheet_name:
            table.save_table_dimensions(sheet_name)

    def on_data_changed(self):
        """Обработчик изменения данных"""
        if not self.is_modified:
            self.is_modified = True
            self.status_label.setText("🔴 Есть несохранённые изменения")
            self.refresh_color_legend()
        
        self.update_size_info()
    
    def update_size_info(self):
        """Обновляет информацию о размере"""
        table = self.get_current_table()
        if table:
            self.size_label.setText(f"Размер: {table.rowCount()}x{table.columnCount()}")
    
    def update_cell_info(self, row, col, prev_row, prev_col):
        """Обновляет информацию о текущей ячейке"""
        if row >= 0 and col >= 0:
            self.cell_info_label.setText(f"Ячейка: [{row+1}, {col+1}]")
    
    def _update_title(self):
        """Обновляет заголовок"""
        if self.current_country_name:
            self.title_label.setText(f"📅 Погодовка - {self.current_country_name}")
        else:
            self.title_label.setText("📅 Погодовка - Все страны")
    
    def get_country_file_path(self, country_id):
        """Возвращает путь к файлу погодовки для страны"""
        self.yearly_dir.mkdir(parents=True, exist_ok=True)
        
        if country_id is not None:
            return self.yearly_dir / f"country_{country_id}.xlsx"
        else:
            if self.current_country_name:
                safe_name = "".join(c for c in self.current_country_name if c.isalnum() or c in ' _-')
                return self.yearly_dir / f"{safe_name}.xlsx"
        return self.yearly_dir / "all_countries.xlsx"

    def set_country(self, country_id, country_name):
        """Устанавливает текущую страну и загружает её погодовку"""
        self.logger.info(f"🔄 set_country: country_id={country_id}, country_name={country_name}")
        self.logger.info(f"   Текущий _loaded_country_id={self._loaded_country_id}")
        
        # Если уже загружена эта страна — ничего не делаем
        if country_id == self._loaded_country_id and country_name == self._loaded_country_name:
            self.logger.info(f"   ⏭️ Страна уже загружена, пропускаем")
            return
        
        # Сохраняем текущую таблицу в кэш перед переключением
        if self._loaded_country_id is not None and self._loaded_country_id in self.loaded_tables_cache:
            self.logger.info(f"   💾 Сохраняем в кэш: {self._loaded_country_id}")
            self._save_to_cache(self._loaded_country_id)
        
        # Сохраняем изменения перед переключением
        if self.is_modified and self.current_file:
            self.logger.info("   💾 Сохраняем изменения перед переключением...")
            self.save_with_formatting(self.current_file)
            self.is_modified = False
        
        self.current_country_id = country_id
        self.current_country_name = country_name
        
        self._update_title()
        
        # Проверяем кэш
        cache_key = country_id
        if cache_key in self.loaded_tables_cache:
            self.logger.info(f"   📂 Восстановление из кэша для {country_name} (ID: {country_id})")
            self._restore_from_cache(self.loaded_tables_cache[cache_key])
            self._loaded_country_id = country_id
            self._loaded_country_name = country_name
            self.current_file = self.loaded_tables_cache[cache_key].get('current_file')
            self.is_modified = False
            self.status_label.setText(f"✅ Восстановлена погодовка для {country_name} (из кэша)")
            self.refresh_color_legend()
            return
        
        # Если нет в кэше — загружаем из файла
        self._load_country_table()

    def _load_country_table(self):
        """Реально загружает погодовку для текущей страны"""
        country_id = self.current_country_id
        country_name = self.current_country_name
        
        if country_id is None and not country_name:
            self.logger.warning("❌ Нет country_id и country_name для загрузки")
            self.create_empty_table()
            return
        
        self.logger.info(f"📂 Загрузка погодовки для {country_name} (ID: {country_id})")
        
        if self._loading_table:
            self.logger.info(f"⏳ Уже идёт загрузка, пропускаем")
            return
        
        self._loading_table = True
        self._loading_country_id = country_id
        self._loading_country_name = country_name
        
        try:
            file_path = self.get_country_file_path(country_id)
            self.logger.info(f"📁 Путь к файлу: {file_path}")
            
            # Сохраняем current_file ДО загрузки
            self.current_file = str(file_path)
            
            if file_path.exists():
                self.logger.info(f"📖 Загрузка из файла: {file_path}")
                self.load_excel_with_formatting(str(file_path))
                self._loaded_country_id = country_id
                self._loaded_country_name = country_name
                self.is_modified = False
                self.status_label.setText(f"✅ Загружена погодовка для {country_name}")
                
                cache_key = country_id
                self._save_to_cache(cache_key)
            else:
                self.logger.info(f"📄 Файл не найден, создаём новую таблицу для {country_name}")
                self.create_empty_table()
                self._loaded_country_id = country_id
                self._loaded_country_name = country_name
                self.is_modified = False
                self.status_label.setText(f"✅ Создана новая погодовка для {country_name}")
                # Сохраняем пустую таблицу в кэш
                cache_key = country_id
                self._save_to_cache(cache_key)
            
            self._update_title()
            self.refresh_color_legend()
            
        except Exception as e:
            self.logger.error(f"❌ Ошибка загрузки: {e}")
            import traceback
            traceback.print_exc()
            self.create_empty_table()
            self.status_label.setText(f"❌ Ошибка загрузки: {str(e)[:50]}")
        finally:
            self._loading_table = False
            self._loading_country_id = None
            self._loading_country_name = None

    def _save_to_cache(self, cache_key):
        """Сохраняет текущую таблицу в кэш"""
        try:
            if cache_key is None:
                return
            
            # Проверяем, есть ли данные для сохранения
            if not self.sheets:
                return
            
            # Получаем актуальный путь к файлу
            file_path = self.get_country_file_path(cache_key)
            file_mtime = 0
            if file_path.exists():
                file_mtime = file_path.stat().st_mtime
            
            # Ограничиваем размер кэша (максимум 10 стран)
            if len(self.loaded_tables_cache) >= 10:
                oldest_key = next(iter(self.loaded_tables_cache))
                del self.loaded_tables_cache[oldest_key]
                self.logger.info(f"   🗑️ Удалён старый кэш: {oldest_key}")
            
            cache_data = {
                'sheets': {},
                'current_file': str(file_path),
                'is_modified': self.is_modified,
                '_file_mtime': file_mtime,
                'country_id': cache_key,
                'country_name': self.current_country_name
            }
            
            for sheet_name, sheet in self.sheets.items():
                table = sheet.table
                
                years = []
                for row in range(table.rowCount()):
                    header = table.verticalHeaderItem(row)
                    year = header.text() if header else ""
                    years.append(year)
                
                sheet_data = {
                    'row_count': table.rowCount(),
                    'col_count': table.columnCount(),
                    'data': [],
                    'headers': [],
                    'column_widths': [],
                    'row_heights': [],
                    'years': years
                }
                
                for col in range(table.columnCount()):
                    header = table.horizontalHeaderItem(col)
                    sheet_data['headers'].append(header.text() if header else f"Колонка {col+1}")
                
                for row in range(table.rowCount()):
                    row_data = []
                    for col in range(table.columnCount()):
                        item = table.item(row, col)
                        if item:
                            cell_data = {
                                'text': item.text(),
                                'formula': item.data(Qt.UserRole + 1) if item.data(Qt.UserRole + 1) else None,
                            }
                            bg_color = item.background().color()
                            if bg_color.isValid() and bg_color.name() != "#ffffff":
                                cell_data['background'] = bg_color.name()
                            fg_color = item.foreground().color()
                            if fg_color.isValid() and fg_color.name() != "#000000":
                                cell_data['foreground'] = fg_color.name()
                            if item.textAlignment() != Qt.AlignLeft | Qt.AlignVCenter:
                                cell_data['alignment'] = item.textAlignment()
                            row_data.append(cell_data)
                        else:
                            row_data.append(None)
                    sheet_data['data'].append(row_data)
                
                for col in range(table.columnCount()):
                    sheet_data['column_widths'].append(table.columnWidth(col))
                
                for row in range(table.rowCount()):
                    sheet_data['row_heights'].append(table.rowHeight(row))
                
                cache_data['sheets'][sheet_name] = sheet_data
            
            self.loaded_tables_cache[cache_key] = cache_data
            self.logger.info(f"   💾 Сохранено в кэш: {cache_key} ({len(cache_data['sheets'])} листов), файл: {file_path}")
            
        except Exception as e:
            self.logger.error(f"Ошибка сохранения в кэш: {e}")

    def _restore_from_cache(self, cache_data):
        """Восстанавливает таблицу из кэша"""
        from PySide6.QtGui import QColor, QBrush
        from PySide6.QtWidgets import QTableWidgetItem
        
        try:
            self.sheet_tabs.clear()
            self.sheets.clear()
            
            for sheet_name, sheet_data in cache_data['sheets'].items():
                from .sheet_widget import SheetWidget
                sheet = SheetWidget(sheet_name)
                table = sheet.table
                
                table.setRowCount(sheet_data['row_count'])
                table.setColumnCount(sheet_data['col_count'])
                
                for col, header_text in enumerate(sheet_data['headers']):
                    if col < table.columnCount():
                        header_item = QTableWidgetItem(header_text)
                        table.setHorizontalHeaderItem(col, header_item)
                
                if 'years' in sheet_data:
                    for row, year in enumerate(sheet_data['years']):
                        if row < table.rowCount():
                            header_item = QTableWidgetItem(str(year))
                            header_item.setBackground(QBrush(QColor("#4a6fa5")))
                            header_item.setForeground(QBrush(QColor("#ffffff")))
                            header_item.setTextAlignment(Qt.AlignCenter)
                            table.setVerticalHeaderItem(row, header_item)
                
                for row, row_data in enumerate(sheet_data['data']):
                    if row >= table.rowCount():
                        table.insertRow(table.rowCount())
                    for col, cell_data in enumerate(row_data):
                        if cell_data and col < table.columnCount():
                            item = QTableWidgetItem(cell_data.get('text', ''))
                            if cell_data.get('formula'):
                                item.setData(Qt.UserRole + 1, cell_data['formula'])
                            if cell_data.get('background'):
                                item.setBackground(QBrush(QColor(cell_data['background'])))
                            if cell_data.get('foreground'):
                                item.setForeground(QBrush(QColor(cell_data['foreground'])))
                            if cell_data.get('alignment'):
                                item.setTextAlignment(cell_data['alignment'])
                            table.setItem(row, col, item)
                            item.setData(Qt.UserRole + 2, True)
                
                for col, width in enumerate(sheet_data['column_widths']):
                    if col < table.columnCount() and width > 0:
                        table.setColumnWidth(col, width)
                
                for row, height in enumerate(sheet_data['row_heights']):
                    if row < table.rowCount() and height > 0:
                        table.setRowHeight(row, height)
                
                self.sheets[sheet_name] = sheet
                self.sheet_tabs.addTab(sheet, sheet_name)
                
                table.dataChanged.connect(self.on_data_changed)
                table.currentCellChanged.connect(self.update_cell_info)
            
            # Восстанавливаем путь к файлу
            self.current_file = cache_data.get('current_file')
            self.is_modified = cache_data.get('is_modified', False)
            
            self.logger.info(f"✅ Восстановлено из кэша: {len(cache_data['sheets'])} листов, файл: {self.current_file}")
            
        except Exception as e:
            self.logger.error(f"Ошибка восстановления из кэша: {e}")
            self.create_empty_table()

    def load_excel_with_formatting(self, file_path):
        """Загружает Excel файл с сохранением форматирования"""
        if not OPENPYXL_AVAILABLE:
            QMessageBox.warning(self, "Предупреждение", 
                               "Библиотека openpyxl не установлена.\n"
                               "Форматирование не будет сохранено.\n"
                               "Установите: pip install openpyxl")
            try:
                import pandas as pd
                df = pd.read_excel(file_path)
                self.load_dataframe(df)
            except Exception as e2:
                QMessageBox.critical(self, "Ошибка", f"Не удалось загрузить файл:\n{e2}")
            return
        
        try:
            from openpyxl.utils import get_column_letter
            from PySide6.QtGui import QColor, QBrush
            from datetime import datetime
            
            wb = load_workbook(file_path, data_only=False)
            self.current_workbook = wb
            
            self.sheet_tabs.clear()
            self.sheets.clear()
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                self.add_new_sheet(sheet_name)
                sheet = self.sheets[sheet_name]
                table = sheet.table
                table.sheet_name = sheet_name
                
                all_rows = list(ws.iter_rows(values_only=False))
                if not all_rows:
                    continue
                
                max_row = len(all_rows)
                max_col = len(all_rows[0]) if all_rows else 0
                
                self.logger.info(f"Загружается лист '{sheet_name}': строк={max_row}, колонок={max_col}")
                
                if max_row <= 1 or max_col <= 0:
                    continue
                
                # Определяем заголовки
                header_row = all_rows[0]
                year_col_index = -1
                headers = []
                
                for col_idx, cell in enumerate(header_row):
                    cell_value = str(cell.value).strip() if cell.value else ""
                    if cell_value == "Год" or cell_value == "Year":
                        year_col_index = col_idx
                        break
                
                if year_col_index == -1:
                    year_col_index = 0
                
                for col_idx, cell in enumerate(header_row):
                    cell_value = str(cell.value).strip() if cell.value else ""
                    if col_idx != year_col_index and cell_value and cell_value != "Год" and cell_value != "Year":
                        headers.append(cell_value)
                
                if not headers:
                    for col_idx in range(max_col):
                        if col_idx != year_col_index:
                            headers.append(f"Колонка_{col_idx}")
                
                table.setRowCount(max_row - 1)
                table.setColumnCount(len(headers))
                
                # Заголовки колонок
                for col_idx, header_text in enumerate(headers):
                    header_item = QTableWidgetItem(header_text)
                    header_item.setTextAlignment(Qt.AlignCenter)
                    header_item.setForeground(QBrush(QColor("#ffffff")))
                    header_item.setBackground(QBrush(QColor("#4a6fa5")))
                    if col_idx < len(header_row):
                        excel_cell = header_row[col_idx + (1 if col_idx >= year_col_index else 0)]
                        self._apply_cell_formatting(header_item, excel_cell)
                    table.setHorizontalHeaderItem(col_idx, header_item)
                
                # Данные
                for row_idx in range(1, max_row):
                    excel_row = all_rows[row_idx]
                    
                    # Год
                    year_value = ""
                    if year_col_index < len(excel_row):
                        year_cell = excel_row[year_col_index]
                        year_value = str(year_cell.value).strip() if year_cell.value else ""
                        header_item = QTableWidgetItem(year_value)
                        header_item.setBackground(QBrush(QColor("#4a6fa5")))
                        header_item.setForeground(QBrush(QColor("#ffffff")))
                        header_item.setTextAlignment(Qt.AlignCenter)
                        table.setVerticalHeaderItem(row_idx - 1, header_item)
                    
                    # Данные по колонкам
                    for col_idx, header_text in enumerate(headers):
                        excel_col_idx = col_idx
                        if excel_col_idx >= year_col_index:
                            excel_col_idx += 1
                        
                        if excel_col_idx < len(excel_row):
                            excel_cell = excel_row[excel_col_idx]
                            value = excel_cell.value
                            
                            item = QTableWidgetItem()
                            item.setForeground(QBrush(QColor("#000000")))
                            
                            # Формулы
                            if isinstance(value, str) and value.startswith('='):
                                item.setData(Qt.UserRole + 1, value)
                                try:
                                    result = table.evaluate_formula(value, row_idx - 1, col_idx)
                                    item.setText(str(result) if result is not None else "#ERROR")
                                    item.setForeground(QBrush(QColor(0, 100, 0)))
                                except Exception as e:
                                    self.logger.debug(f"Ошибка формулы {value}: {e}")
                                    item.setText("#ERROR")
                            else:
                                if value is not None:
                                    if isinstance(value, (int, float)):
                                        if isinstance(value, float):
                                            item.setText(f"{value:.2f}")
                                        else:
                                            item.setText(str(value))
                                        item.setData(Qt.UserRole, value)
                                    elif isinstance(value, datetime):
                                        item.setText(value.strftime("%Y-%m-%d"))
                                    else:
                                        item.setText(str(value))
                                else:
                                    item.setText("")
                            
                            self._apply_cell_formatting(item, excel_cell)
                            
                            # Пустые ячейки
                            if not value or value == "" or value == "-" or value == "—":
                                item.setBackground(QBrush(QColor("#6c757d")))
                                item.setForeground(QBrush(QColor("#ffffff")))
                            
                            table.setItem(row_idx - 1, col_idx, item)
                            item.setData(Qt.UserRole + 2, True)
                        else:
                            item = QTableWidgetItem("")
                            item.setBackground(QBrush(QColor("#6c757d")))
                            item.setForeground(QBrush(QColor("#ffffff")))
                            table.setItem(row_idx - 1, col_idx, item)
                            item.setData(Qt.UserRole + 2, True)
                
                # Восстанавливаем размеры
                table.restore_table_dimensions(sheet_name)
                
                # Ширина колонок из файла
                for col_idx in range(len(headers)):
                    try:
                        excel_col_idx = col_idx
                        if excel_col_idx >= year_col_index:
                            excel_col_idx += 1
                        col_letter = get_column_letter(excel_col_idx + 1)
                        col_width = ws.column_dimensions[col_letter].width
                        if col_width and col_width > 0:
                            pixel_width = int(col_width * 7)
                            if pixel_width > 0:
                                table.setColumnWidth(col_idx, pixel_width)
                    except:
                        pass
                
                # Высота строк из файла
                for row_idx in range(max_row - 1):
                    try:
                        row_height = ws.row_dimensions[row_idx + 2].height
                        if row_height and row_height > 0:
                            pixel_height = int(row_height)
                            if pixel_height > 0:
                                table.setRowHeight(row_idx, pixel_height)
                    except:
                        pass
                
                # Объединённые ячейки
                for merged_range in ws.merged_cells.ranges:
                    try:
                        min_col = merged_range.min_col
                        max_col = merged_range.max_col
                        min_row = merged_range.min_row
                        max_row_merged = merged_range.max_row
                        
                        if year_col_index + 1 >= min_col and year_col_index + 1 <= max_col:
                            continue
                        
                        top = min_row - 2
                        left = min_col - 1
                        if left > year_col_index:
                            left -= 1
                        
                        rows = max_row_merged - min_row + 1
                        cols = max_col - min_col + 1
                        
                        if top >= 0 and left >= 0 and top + rows <= table.rowCount() and left + cols <= table.columnCount():
                            table.setSpan(top, left, rows, cols)
                            merge_info = {'top': top, 'left': left, 'rows': rows, 'cols': cols}
                            table.merged_cells.append(merge_info)
                    except Exception as e:
                        self.logger.debug(f"Ошибка восстановления объединения: {e}")
                
                self.logger.info(f"✅ Загружен лист '{sheet_name}': {max_row-1} строк, {len(headers)} колонок")
            
            self.logger.info(f"✅ Всего загружено листов: {len(wb.sheetnames)}")
            self.is_modified = False
            self.refresh_color_legend()
            self.restore_tab_order()
            
        except Exception as e:
            self.logger.error(f"❌ Ошибка загрузки: {e}")
            import traceback
            traceback.print_exc()
            try:
                import pandas as pd
                df = pd.read_excel(file_path)
                self.load_dataframe(df)
                QMessageBox.information(self, "Информация", "Файл загружен без форматирования")
            except Exception as e2:
                QMessageBox.critical(self, "Ошибка", f"Не удалось загрузить файл:\n{e2}")

    def _apply_cell_formatting(self, item, cell):
        """Применяет форматирование из ячейки Excel к элементу таблицы"""
        try:
            if cell.fill and hasattr(cell.fill, 'fgColor') and cell.fill.fgColor:
                try:
                    rgb_value = cell.fill.fgColor.rgb
                    color_hex = None
                    
                    if rgb_value:
                        if isinstance(rgb_value, str):
                            color_hex = rgb_value
                        elif hasattr(rgb_value, 'value'):
                            color_hex = str(rgb_value.value)
                        elif hasattr(rgb_value, 'rgb'):
                            color_hex = str(rgb_value.rgb)
                        else:
                            color_hex = str(rgb_value)
                        
                        if color_hex and color_hex != '00000000':
                            if len(color_hex) == 8:
                                color_hex = color_hex[2:]
                            color_hex = color_hex.lstrip('#')
                            if len(color_hex) == 6:
                                try:
                                    color = QColor(f"#{color_hex}")
                                    if color.isValid():
                                        item.setBackground(QBrush(color))
                                except:
                                    pass
                except Exception as e:
                    self.logger.debug(f"Ошибка при обработке цвета заливки: {e}")
            
            if cell.font:
                font = QFont()
                if cell.font.name:
                    font.setFamily(cell.font.name)
                if cell.font.sz:
                    try:
                        font.setPointSize(int(cell.font.sz))
                    except:
                        pass
                if cell.font.b:
                    font.setBold(True)
                if cell.font.i:
                    font.setItalic(True)
                if cell.font.u:
                    font.setUnderline(True)
                item.setFont(font)
            
            if cell.alignment:
                horizontal = cell.alignment.horizontal
                vertical = cell.alignment.vertical
                
                align = 0
                
                if horizontal == 'right':
                    align |= Qt.AlignRight
                elif horizontal == 'center':
                    align |= Qt.AlignCenter
                elif horizontal == 'justify':
                    align |= Qt.AlignJustify
                else:
                    align |= Qt.AlignLeft
                
                if vertical == 'top':
                    align |= Qt.AlignTop
                elif vertical == 'bottom':
                    align |= Qt.AlignBottom
                else:
                    align |= Qt.AlignVCenter
                
                item.setTextAlignment(align)
            
            if cell.font and cell.font.color and cell.font.color.rgb:
                try:
                    rgb_value = cell.font.color.rgb
                    color_hex = None
                    
                    if rgb_value:
                        if isinstance(rgb_value, str):
                            color_hex = rgb_value
                        elif hasattr(rgb_value, 'value'):
                            color_hex = str(rgb_value.value)
                        elif hasattr(rgb_value, 'rgb'):
                            color_hex = str(rgb_value.rgb)
                        else:
                            color_hex = str(rgb_value)
                        
                        if color_hex and color_hex != '00000000':
                            if len(color_hex) == 8:
                                color_hex = color_hex[2:]
                            color_hex = color_hex.lstrip('#')
                            if len(color_hex) == 6:
                                try:
                                    color = QColor(f"#{color_hex}")
                                    if color.isValid():
                                        item.setForeground(QBrush(color))
                                except:
                                    pass
                except Exception as e:
                    self.logger.debug(f"Ошибка при обработке цвета текста: {e}")
            
            if cell.border:
                table = self.get_current_table()
                if table:
                    existing_borders = item.data(Qt.UserRole + 2)
                    if existing_borders and isinstance(existing_borders, dict):
                        borders = existing_borders.copy()
                    else:
                        borders = {}
                    
                    if cell.border.top:
                        borders['top'] = self._get_border_type_from_side(cell.border.top)
                    elif 'top' not in borders:
                        borders['top'] = table.BORDER_NONE
                    
                    if cell.border.bottom:
                        borders['bottom'] = self._get_border_type_from_side(cell.border.bottom)
                    elif 'bottom' not in borders:
                        borders['bottom'] = table.BORDER_NONE
                    
                    if cell.border.left:
                        borders['left'] = self._get_border_type_from_side(cell.border.left)
                    elif 'left' not in borders:
                        borders['left'] = table.BORDER_NONE
                    
                    if cell.border.right:
                        borders['right'] = self._get_border_type_from_side(cell.border.right)
                    elif 'right' not in borders:
                        borders['right'] = table.BORDER_NONE
                    
                    item.setData(Qt.UserRole + 2, borders)
                
        except Exception as e:
            self.logger.error(f"Ошибка при применении форматирования: {e}")
    
    def _get_border_type_from_side(self, side):
        """Преобразует стиль границы Excel в тип границы таблицы"""
        table = self.get_current_table()
        if not table:
            return 1
        
        if side.style == 'thin':
            return table.BORDER_THIN
        elif side.style == 'medium':
            return table.BORDER_MEDIUM
        elif side.style == 'thick':
            return table.BORDER_THICK
        elif side.style == 'double':
            return table.BORDER_DOUBLE
        else:
            return table.BORDER_THIN

    def open_file(self):
        """Открывает файл Excel"""
        self.yearly_dir.mkdir(parents=True, exist_ok=True)
        
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Открыть файл",
            str(self.yearly_dir),
            "Excel files (*.xlsx *.xls);;CSV files (*.csv);;All files (*.*)"
        )
        
        if file_path:
            try:
                self.load_excel_with_formatting(file_path)
                self.current_file = file_path
                self.is_modified = False
                self.status_label.setText(f"Загружен: {os.path.basename(file_path)}")
                self._update_country_from_filename(os.path.basename(file_path))
                self.refresh_color_legend()
            except Exception as e:
                error_msg = str(e)
                if "xlrd" in error_msg:
                    QMessageBox.critical(self, "Ошибка", 
                        f"Для работы с файлами .xls требуется библиотека xlrd.\n\n"
                        f"Установите: pip install xlrd\n\n"
                        f"Или сохраните файл в формате .xlsx")
                else:
                    QMessageBox.critical(self, "Ошибка", f"Не удалось загрузить файл:\n{error_msg}")

    def _update_country_from_filename(self, filename):
        """Обновляет страну из имени файла"""
        name = os.path.splitext(filename)[0]
        if name.startswith('country_'):
            try:
                self.current_country_id = int(name.replace('country_', ''))
                self.current_country_name = f"Страна {self.current_country_id}"
                self.title_label.setText(f"📅 Погодовка - {self.current_country_name}")
            except:
                pass
        else:
            self.current_country_id = None
            self.current_country_name = name
            self.title_label.setText(f"📅 Погодовка - {self.current_country_name}")

    def save_file(self):
        """Сохраняет таблицу в текущий файл (перезаписывает)"""
        if not self.current_file:
            self.save_file_as()
            return
        
        import shutil
        backup = self.current_file + ".backup"
        if os.path.exists(self.current_file):
            try:
                shutil.copy2(self.current_file, backup)
            except:
                pass
        
        try:
            self.save_with_formatting(self.current_file)
            self.is_modified = False
            self.status_label.setText(f"✅ Сохранено: {os.path.basename(self.current_file)}")
            
            # Обновляем кэш
            cache_key = self.current_country_id
            if cache_key:
                self._save_to_cache(cache_key)
            
            if os.path.exists(backup):
                try:
                    os.remove(backup)
                except:
                    pass
        except Exception as e:
            self.logger.error(f"Ошибка сохранения: {e}")
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{e}")
            
            if os.path.exists(backup):
                try:
                    shutil.copy2(backup, self.current_file)
                    self.status_label.setText("⚠️ Восстановлено из резервной копии")
                except:
                    pass

    def save_file_as(self):
        """Сохраняет таблицу в новый файл (с выбором имени)"""
        default = f"country_{self.current_country_id}.xlsx" if self.current_country_id else f"{self.current_country_name or 'all'}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить файл",
            str(self.yearly_dir / default),
            "Excel files (*.xlsx)"
        )
        
        if file_path:
            if not file_path.endswith('.xlsx'):
                file_path += '.xlsx'
            
            Path(file_path).parent.mkdir(parents=True, exist_ok=True)
            
            self.save_with_formatting(file_path)
            self.current_file = file_path
            self.is_modified = False
            self.status_label.setText(f"✅ Сохранено: {os.path.basename(file_path)}")
            
            # Обновляем кэш
            cache_key = self.current_country_id
            if cache_key:
                self._save_to_cache(cache_key)

    def save_with_formatting(self, file_path):
        """Сохраняет таблицу в Excel с сохранением форматирования"""
        if not OPENPYXL_AVAILABLE:
            self.save_to_file(file_path)
            return
        
        Path(file_path).parent.mkdir(parents=True, exist_ok=True)
        
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
            
            for sheet_name, sheet in self.sheets.items():
                ws = wb.create_sheet(sheet_name)
                table = sheet.table
                table.sheet_name = sheet_name
                
                year_cell = ws.cell(row=1, column=1, value="Год")
                year_cell.font = Font(bold=True, color="FFFFFF")
                year_cell.fill = PatternFill(start_color="4a6fa5", end_color="4a6fa5", fill_type="solid")
                year_cell.alignment = Alignment(horizontal="center", vertical="center")
                
                for col in range(table.columnCount()):
                    header = table.horizontalHeaderItem(col)
                    header_text = header.text() if header else f"Колонка {col+1}"
                    cell = ws.cell(row=1, column=col+2, value=header_text)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="4a6fa5", end_color="4a6fa5", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                
                for row in range(table.rowCount()):
                    year_header = table.verticalHeaderItem(row)
                    year_value = year_header.text() if year_header else ""
                    year_cell = ws.cell(row=row+2, column=1, value=year_value)
                    year_cell.fill = PatternFill(start_color="4a6fa5", end_color="4a6fa5", fill_type="solid")
                    year_cell.font = Font(bold=True, color="FFFFFF")
                    year_cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                    for col in range(table.columnCount()):
                        item = table.item(row, col)
                        cell = ws.cell(row=row+2, column=col+2)
                        
                        if item:
                            formula = item.data(Qt.UserRole + 1)
                            if formula:
                                cell.value = formula
                            else:
                                cell.value = item.text()
                            
                            bg_color = item.background().color()
                            if bg_color.isValid() and bg_color.name() != "#ffffff":
                                color_hex = bg_color.name().lstrip('#')
                                if len(color_hex) == 6:
                                    cell.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
                            
                            fg_color = item.foreground().color()
                            if fg_color.isValid() and fg_color.name() != "#000000":
                                color_hex = fg_color.name().lstrip('#')
                                if len(color_hex) == 6:
                                    cell.font = Font(color=color_hex)
                            
                            self._apply_item_alignment_to_cell(item, cell)
                            self._apply_item_borders_to_cell(item, cell)
                        else:
                            cell.fill = PatternFill(start_color="9a9a73", end_color="9a9a73", fill_type="solid")
                
                ws.column_dimensions['A'].width = 12
                
                for col in range(table.columnCount()):
                    col_width = table.columnWidth(col)
                    if col_width > 0:
                        col_letter = get_column_letter(col + 2)
                        ws.column_dimensions[col_letter].width = col_width / 7
                
                for row in range(table.rowCount()):
                    row_height = table.rowHeight(row)
                    if row_height > 0:
                        ws.row_dimensions[row + 2].height = row_height
                
                for merge in table.merged_cells:
                    try:
                        top = merge['top'] + 2
                        left = merge['left'] + 2
                        bottom = top + merge['rows'] - 1
                        right = left + merge['cols'] - 1
                        
                        start_cell = f"{get_column_letter(left)}{top}"
                        end_cell = f"{get_column_letter(right)}{bottom}"
                        
                        ws.merge_cells(f"{start_cell}:{end_cell}")
                    except Exception as e:
                        self.logger.debug(f"Ошибка сохранения объединения: {e}")
                
                table.save_column_widths(sheet_name)
            
            wb.save(file_path)
            self.current_workbook = wb
            self.is_modified = False
            self.status_label.setText(f"✅ Сохранено: {os.path.basename(file_path)}")
            
        except Exception as e:
            self.logger.error(f"Ошибка сохранения: {e}")
            import traceback
            traceback.print_exc()
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{e}")

    def _apply_item_alignment_to_cell(self, item, cell):
        """Применяет выравнивание из QTableWidgetItem к Excel ячейке"""
        from openpyxl.styles import Alignment
        align = int(item.textAlignment())
        # Горизонтальное выравнивание (маска горизонтали)
        h = align & int(Qt.AlignHorizontal_Mask)
        if h == int(Qt.AlignRight):
            horizontal = 'right'
        elif h == int(Qt.AlignHCenter):
            horizontal = 'center'
        elif h == int(Qt.AlignJustify):
            horizontal = 'justify'
        else:
            horizontal = 'left'
        # Вертикальное выравнивание (маска вертикали)
        v = align & int(Qt.AlignVertical_Mask)
        if v == int(Qt.AlignTop):
            vertical = 'top'
        elif v == int(Qt.AlignBottom):
            vertical = 'bottom'
        else:
            vertical = 'center'
        try:
            cell.alignment = Alignment(horizontal=horizontal, vertical=vertical)
        except Exception as e:
            self.logger.debug(f"Ошибка установки выравнивания: {e}")

    def _apply_item_borders_to_cell(self, item, cell):
        """Применяет границы из QTableWidgetItem к Excel ячейке"""
        from openpyxl.styles import Border, Side
        
        table = self.get_current_table()
        if not table:
            return
        
        borders = item.data(Qt.UserRole + 2)
        if not borders or not isinstance(borders, dict):
            return
        
        border_obj = Border()
        
        def get_side(border_type):
            if border_type == table.BORDER_THIN:
                return Side(style='thin', color='000000')
            elif border_type == table.BORDER_MEDIUM:
                return Side(style='medium', color='000000')
            elif border_type == table.BORDER_THICK:
                return Side(style='thick', color='000000')
            elif border_type == table.BORDER_DOUBLE:
                return Side(style='double', color='000000')
            else:
                return None
        
        if borders.get('top', 0) != 0:
            side = get_side(borders['top'])
            if side:
                border_obj.top = side
        if borders.get('bottom', 0) != 0:
            side = get_side(borders['bottom'])
            if side:
                border_obj.bottom = side
        if borders.get('left', 0) != 0:
            side = get_side(borders['left'])
            if side:
                border_obj.left = side
        if borders.get('right', 0) != 0:
            side = get_side(borders['right'])
            if side:
                border_obj.right = side
        
        if border_obj.top or border_obj.bottom or border_obj.left or border_obj.right:
            cell.border = border_obj

    def save_to_file(self, file_path):
        """Сохраняет в файл без pandas"""
        try:
            from openpyxl import Workbook
            
            table = self.get_current_table()
            if not table:
                return
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Лист1"
            
            for col in range(table.columnCount()):
                header = table.horizontalHeaderItem(col)
                header_text = header.text() if header else f"Col{col+1}"
                ws.cell(row=1, column=col+1, value=header_text)
            
            for row in range(table.rowCount()):
                for col in range(table.columnCount()):
                    item = table.item(row, col)
                    value = item.text() if item else ""
                    ws.cell(row=row+2, column=col+1, value=value)
            
            wb.save(file_path)
            self.is_modified = False
            
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить:\n{e}")

    def save_current_table(self):
        """Сохраняет текущую таблицу в файл (перезаписывает существующий)"""
        if not self.current_file:
            default_name = f"country_{self.current_country_id}.xlsx" if self.current_country_id else f"{self.current_country_name or 'all'}.xlsx"
            file_path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить файл",
                str(self.yearly_dir / default_name),
                "Excel files (*.xlsx)"
            )
            if not file_path:
                return False
            if not file_path.endswith('.xlsx'):
                file_path += '.xlsx'
            self.current_file = file_path
        
        self.save_with_formatting(self.current_file)
        self.is_modified = False
        self.status_label.setText(f"✅ Сохранено: {os.path.basename(self.current_file)}")
        
        cache_key = self.current_country_id
        if cache_key:
            self._save_to_cache(cache_key)
        
        return True

    def create_empty_table(self):
        """Создает пустую таблицу"""
        self.sheet_tabs.clear()
        self.sheets.clear()
        self.add_new_sheet("Лист1")
        
        if self.sheet_tabs.count() > 0:
            current_sheet = self.sheet_tabs.currentWidget()
            if current_sheet and hasattr(current_sheet, 'table'):
                current_sheet.table.setAlternatingRowColors(False)
        
        self.is_modified = False
        self.current_workbook = None
        self.update_size_info()
        self.refresh_color_legend()
        
        if self.current_country_id is not None:
            settings_key = f'tab_order_country_{self.current_country_id}'
        else:
            settings_key = f'tab_order_all'
        self.settings.remove(settings_key)

    def create_new_table(self):
        """Создаёт новую таблицу"""
        self.create_empty_table()
        
        if self.current_country_id is not None:
            file_path = self.get_country_file_path(self.current_country_id)
        else:
            file_path = self.get_country_file_path(None)
        
        file_path.parent.mkdir(parents=True, exist_ok=True)
        self.current_file = str(file_path)
        
        self.status_label.setText("Создана новая таблица")
        self.refresh_color_legend()
        
        cache_key = self.current_country_id
        if cache_key:
            self._save_to_cache(cache_key)
        QTimer.singleShot(0, self._restore_sheet_order)
    
    def refresh_current(self):
        """Обновляет текущую погодовку (принудительная перезагрузка)"""
        if self.current_country_id is not None:
            cache_key = self.current_country_id
            if cache_key in self.loaded_tables_cache:
                del self.loaded_tables_cache[cache_key]
            
            self._loaded_country_id = None
            self._loaded_country_name = None
            
            self._load_country_table()
        elif self.current_country_name:
            self._load_country_table()

    def auto_save(self):
        """Автосохранение (перезаписывает существующий файл)"""
        if self.auto_save_enabled and self.is_modified and self.current_file:
            try:
                self.save_with_formatting(self.current_file)
                self.is_modified = False
                self.status_label.setText("✅ Автосохранение выполнено")
                
                cache_key = self.current_country_id
                if cache_key:
                    self._save_to_cache(cache_key)
                
                from PySide6.QtCore import QTimer
                QTimer.singleShot(2000, lambda: self.status_label.setText("Готов"))
            except Exception as e:
                self.logger.error(f"Ошибка автосохранения: {e}")

    def save_all_sheets(self):
        """Сохраняет все листы в файл (перезаписывает существующий)"""
        if not self.current_file:
            default_name = f"country_{self.current_country_id}.xlsx" if self.current_country_id else f"{self.current_country_name or 'all'}.xlsx"
            file_path, _ = QFileDialog.getSaveFileName(
                self, "Сохранить файл",
                str(self.yearly_dir / default),
                "Excel files (*.xlsx)"
            )
            if not file_path:
                return False
            if not file_path.endswith('.xlsx'):
                file_path += '.xlsx'
            self.current_file = file_path
        try:
            # Сначала сохраняем размеры текущей таблицы (ширины колонок + высоты строк)
            current_widget = self.sheet_tabs.currentWidget()
            if current_widget and hasattr(current_widget, 'table'):
                table = current_widget.table
                sheet_name = self.sheet_tabs.tabText(self.sheet_tabs.currentIndex())
                table.sheet_name = sheet_name
                table.save_table_dimensions(sheet_name)
            # Сохраняем в Excel (включая высоты строк и выравнивание ячеек)
            self.save_with_formatting(self.current_file)
            self.is_modified = False
            self.status_label.setText(f"✅ Все изменения сохранены: {os.path.basename(self.current_file)}")
            cache_key = self.current_country_id
            if cache_key:
                self._save_to_cache(cache_key)
            self.logger.info(f"💾 Кэш обновлён для {cache_key}")
            return True
        except Exception as e:
            self.logger.error(f"Ошибка сохранения: {e}")
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{e}")
            return False

    def check_save_on_close(self):
        """Проверяет необходимость сохранения"""
        if self.is_modified:
            reply = QMessageBox.question(
                self,
                "Несохраненные изменения",
                "Сохранить изменения перед закрытием?",
                QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel
            )
            if reply == QMessageBox.Cancel:
                return False
            if reply == QMessageBox.Yes:
                self.save_current_table()
        return True

    def rename_current_sheet(self):
        """Переименовывает текущую вкладку"""
        current_index = self.sheet_tabs.currentIndex()
        if current_index < 0:
            QMessageBox.warning(self, "Предупреждение", "Нет активной вкладки")
            return
        
        old_name = self.sheet_tabs.tabText(current_index)
        
        new_name, ok = QInputDialog.getText(
            self, 
            "Переименовать вкладку",
            "Введите новое название:",
            text=old_name
        )
        
        if ok and new_name and new_name.strip():
            new_name = new_name.strip()
            
            for i in range(self.sheet_tabs.count()):
                if i != current_index and self.sheet_tabs.tabText(i) == new_name:
                    QMessageBox.warning(self, "Предупреждение", f"Вкладка '{new_name}' уже существует")
                    return
            
            self.sheet_tabs.setTabText(current_index, new_name)
            
            sheet = self.sheets.pop(old_name)
            sheet.sheet_name = new_name
            self.sheets[new_name] = sheet
            
            self.is_modified = True
            self.status_label.setText(f"✅ Вкладка переименована: {old_name} → {new_name}")
            self.refresh_color_legend()
            self.save_tab_order()

    def show_tab_context_menu(self, position):
        """Показывает контекстное меню для вкладок"""
        tab_index = self.sheet_tabs.tabBar().tabAt(position)
        if tab_index < 0:
            return
        
        menu = QMenu()
        
        rename_action = QAction("✏️ Переименовать", self)
        rename_action.triggered.connect(self.rename_current_sheet)
        menu.addAction(rename_action)
        
        menu.exec(self.sheet_tabs.mapToGlobal(position))

    def show_borders_menu(self):
        """Показывает меню выбора границ как в Excel"""
        table = self.get_current_table()
        if not table:
            QMessageBox.warning(self, "Предупреждение", "Нет активной таблицы")
            return
        
        selected_items = table.selectedItems()
        if not selected_items:
            QMessageBox.warning(self, "Предупреждение", "Выделите ячейки для применения границ")
            return
        
        menu = QMenu(self)
        menu.setWindowTitle("Границы")
        
        bottom_action = QAction("⬇ Нижняя граница", menu)
        bottom_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'bottom'))
        menu.addAction(bottom_action)
        
        top_action = QAction("⬆ Верхняя граница", menu)
        top_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'top'))
        menu.addAction(top_action)
        
        left_action = QAction("⬅ Левая граница", menu)
        left_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'left'))
        menu.addAction(left_action)
        
        right_action = QAction("➡ Правая граница", menu)
        right_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'right'))
        menu.addAction(right_action)
        
        menu.addSeparator()
        
        none_action = QAction("✖ Нет границы", menu)
        none_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_NONE, 'all'))
        menu.addAction(none_action)
        
        all_action = QAction("⬚ Все границы", menu)
        all_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'all'))
        menu.addAction(all_action)
        
        outer_action = QAction("⬚ Внешние границы", menu)
        outer_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'outer'))
        menu.addAction(outer_action)
        
        menu.addSeparator()
        
        thick_outer_action = QAction("━ Толстые внешние границы", menu)
        thick_outer_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THICK, 'outer'))
        menu.addAction(thick_outer_action)
        
        menu.addSeparator()
        
        double_bottom_action = QAction("═ Сдвоенная нижняя граница", menu)
        double_bottom_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_DOUBLE, 'bottom'))
        menu.addAction(double_bottom_action)
        
        thick_bottom_action = QAction("━ Толстая нижняя граница", menu)
        thick_bottom_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THICK, 'bottom'))
        menu.addAction(thick_bottom_action)
        
        menu.addSeparator()
        
        top_bottom_action = QAction("⬆⬇ Верхняя и нижняя границы", menu)
        top_bottom_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'top_bottom'))
        menu.addAction(top_bottom_action)
        
        top_thick_bottom_action = QAction("⬆ Верхняя + ━ толстая нижняя", menu)
        top_thick_bottom_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'top_thick_bottom'))
        menu.addAction(top_thick_bottom_action)
        
        top_double_bottom_action = QAction("⬆ Верхняя + ═ сдвоенная нижняя", menu)
        top_double_bottom_action.triggered.connect(lambda: self._apply_border_to_selection(table, table.BORDER_THIN, 'top_double_bottom'))
        menu.addAction(top_double_bottom_action)
        
        menu.exec(self.mapToGlobal(self.sender().pos()))

    def _apply_border_to_selection(self, table, border_type, side):
        """Применяет границы к выделенным ячейкам"""
        table._save_state()
        
        selected_ranges = table.selectedRanges()
        if not selected_ranges:
            return
        
        cells = set()
        for item in table.selectedItems():
            cells.add((item.row(), item.column()))
        
        if side == 'outer':
            for sel in selected_ranges:
                top = sel.topRow()
                bottom = sel.bottomRow()
                left = sel.leftColumn()
                right = sel.rightColumn()
                
                for col in range(left, right + 1):
                    item = table.item(top, col)
                    if not item:
                        item = QTableWidgetItem()
                        table.setItem(top, col, item)
                    table._apply_border_to_item(item, border_type, 'top', top, col)
                
                for col in range(left, right + 1):
                    item = table.item(bottom, col)
                    if not item:
                        item = QTableWidgetItem()
                        table.setItem(bottom, col, item)
                    table._apply_border_to_item(item, border_type, 'bottom', bottom, col)
                
                for row in range(top, bottom + 1):
                    item = table.item(row, left)
                    if not item:
                        item = QTableWidgetItem()
                        table.setItem(row, left, item)
                    table._apply_border_to_item(item, border_type, 'left', row, left)
                
                for row in range(top, bottom + 1):
                    item = table.item(row, right)
                    if not item:
                        item = QTableWidgetItem()
                        table.setItem(row, right, item)
                    table._apply_border_to_item(item, border_type, 'right', row, right)
            
            table.viewport().update()
            table.dataChanged.emit()
            return
        
        elif side == 'all':
            sides_to_apply = ['top', 'bottom', 'left', 'right']
            for side_name in sides_to_apply:
                for row, col in cells:
                    item = table.item(row, col)
                    if not item:
                        item = QTableWidgetItem()
                        table.setItem(row, col, item)
                    table._apply_border_to_item(item, border_type, side_name, row, col)
        
        elif side == 'top_bottom':
            sides_to_apply = ['top', 'bottom']
            for side_name in sides_to_apply:
                for row, col in cells:
                    item = table.item(row, col)
                    if not item:
                        item = QTableWidgetItem()
                        table.setItem(row, col, item)
                    table._apply_border_to_item(item, border_type, side_name, row, col)
        
        elif side == 'top_thick_bottom':
            for row, col in cells:
                item = table.item(row, col)
                if not item:
                    item = QTableWidgetItem()
                    table.setItem(row, col, item)
                table._apply_border_to_item(item, border_type, 'top', row, col)
                table._apply_border_to_item(item, table.BORDER_THICK, 'bottom', row, col)
        
        elif side == 'top_double_bottom':
            for row, col in cells:
                item = table.item(row, col)
                if not item:
                    item = QTableWidgetItem()
                    table.setItem(row, col, item)
                table._apply_border_to_item(item, border_type, 'top', row, col)
                table._apply_border_to_item(item, table.BORDER_DOUBLE, 'bottom', row, col)
        
        else:
            for row, col in cells:
                item = table.item(row, col)
                if not item:
                    item = QTableWidgetItem()
                    table.setItem(row, col, item)
                table._apply_border_to_item(item, border_type, side, row, col)
        
        table.viewport().update()
        table.dataChanged.emit()

    def auto_fit_columns(self):
        """Автоширина: колонки по содержимому, строки 5px, текст по центру"""
        from PySide6.QtWidgets import QTableView, QHeaderView
        current = self.sheet_tabs.currentWidget() if hasattr(self, 'sheet_tabs') else None
        if current is None:
            return
        table = current.findChild(QTableView)
        if table is None:
            return
        table.setUpdatesEnabled(False)
        try:
            header = table.horizontalHeader()
            # Сначала сбрасываем режим, чтобы resize сработал для всех колонок
            for i in range(header.count()):
                header.setSectionResizeMode(i, QHeaderView.Interactive)
            # Ширина колонок по содержимому
            table.resizeColumnsToContents()
            # Высота ВСЕХ строк = 5px
            for row in range(table.rowCount()):
                table.setRowHeight(row, 5)
            # Текст ВСЕХ ячеек по центру
            for row in range(table.rowCount()):
                for col in range(table.columnCount()):
                    item = table.item(row, col)
                    if item:
                        item.setTextAlignment(Qt.AlignCenter)
        finally:
            table.setUpdatesEnabled(True)
        # Помечаем как изменённое, чтобы SAVE ALL сохранил изменения
        self.is_modified = True
        # Сохраняем размеры (ширины + высоты) в настройки сразу
        sheet_name = self.sheet_tabs.tabText(self.sheet_tabs.currentIndex())
        table.sheet_name = sheet_name
        table.save_table_dimensions(sheet_name)
        # Обновляем кэш, чтобы после переключения/перезапуска всё восстановилось
        if self.current_country_id is not None:
            self._save_to_cache(self.current_country_id)
        if hasattr(self, 'status_label'):
            self.status_label.setText("✅ Автоширина: колонки по содержимому, строки 5px, текст по центру")

    def import_from_ucoin(self):
        """Импортирует таблицу из открытой вкладки браузера uCoin.net"""
        from PySide6.QtWidgets import QProgressDialog
        from gui.widgets.yearly_table.ucoin_parser import UCoinParser, create_yearly_table_from_parser
        from PySide6.QtGui import QColor
        
        main_window = None
        for widget in QApplication.topLevelWidgets():
            if widget.__class__.__name__ == 'MainWindow':
                main_window = widget
                break
        
        if not main_window:
            QMessageBox.warning(self, "Ошибка", "Не найдено главное окно")
            return
        
        browser_tab = None
        if hasattr(main_window, 'center_tab_widget'):
            for i in range(main_window.center_tab_widget.count()):
                tab_text = main_window.center_tab_widget.tabText(i)
                if "Браузер" in tab_text:
                    browser_tab = main_window.center_tab_widget.widget(i)
                    break
        
        if not browser_tab:
            QMessageBox.warning(self, "Ошибка", "Вкладка браузера не найдена")
            return
        
        if hasattr(browser_tab, 'get_current_web_view'):
            web_view = browser_tab.get_current_web_view()
        elif hasattr(browser_tab, 'web_view'):
            web_view = browser_tab.web_view
        else:
            QMessageBox.warning(self, "Ошибка", "Не удалось получить WebView")
            return
        
        if not web_view:
            QMessageBox.warning(self, "Ошибка", "Нет активной страницы в браузере")
            return
        
        current_url = web_view.url().toString()
        if 'ucoin.net' not in current_url:
            reply = QMessageBox.question(
                self, 
                "Подтверждение",
                f"Текущая страница не похожа на uCoin.net.\n"
                f"URL: {current_url[:80]}...\n\n"
                f"Продолжить парсинг?",
                QMessageBox.Yes | QMessageBox.No
            )
            if reply != QMessageBox.Yes:
                return
        
        progress = QProgressDialog("Парсинг страницы uCoin.net...", "Отмена", 0, 100, self)
        progress.setWindowModality(Qt.WindowModal)
        progress.setAutoClose(True)
        progress.setMinimumDuration(500)
        progress.setValue(0)
        progress.show()
        
        parser = UCoinParser(web_view)
        
        def update_progress(value, message):
            if progress.wasCanceled():
                return
            progress.setValue(value)
            progress.setLabelText(message)
            QApplication.processEvents()
        
        data, page_title = parser.parse_current_page(update_progress)
        
        if not data:
            progress.close()
            QMessageBox.warning(self, "Ошибка", "Не удалось распарсить страницу.\n\n"
                               "Убедитесь, что на странице открыта таблица монет uCoin.net\n"
                               "с годами и номиналами.")
            return
        
        if page_title:
            country_name = re.sub(r'[<>:"/\\|?*]', '', page_title).strip()
            if len(country_name) > 50:
                country_name = country_name[:50]
        else:
            country_name = self.current_country_name or "Импорт"
        
        headers_from_parser = getattr(parser, 'headers', None)
        if headers_from_parser:
            self.logger.info(f"📋 Заголовки из таблицы: {headers_from_parser}")
        
        table_data = create_yearly_table_from_parser(data, country_name, headers_from_parser)
        
        if not table_data:
            progress.close()
            QMessageBox.warning(self, "Ошибка", "Не найдены данные о монетах")
            return
        
        progress.setLabelText("Создание таблицы...")
        QApplication.processEvents()
        
        sheet_name = country_name
        base_name = sheet_name
        counter = 1
        while sheet_name in self.sheets:
            sheet_name = f"{base_name}_{counter}"
            counter += 1
        
        self.add_new_sheet(sheet_name)
        
        sheet = self.sheets.get(sheet_name)
        if not sheet:
            progress.close()
            QMessageBox.warning(self, "Ошибка", "Не удалось создать лист")
            return
        
        table = sheet.table
        
        headers = table_data['headers']
        table_rows = table_data['rows']
        years = table_data['years']
        
        table.setRowCount(len(table_rows))
        table.setColumnCount(len(headers))
        
        for col, denom in enumerate(headers):
            header_item = QTableWidgetItem(denom)
            header_item.setTextAlignment(Qt.AlignCenter)
            table.setHorizontalHeaderItem(col, header_item)
        
        for row, year in enumerate(years):
            header_item = QTableWidgetItem(str(year))
            header_item.setTextAlignment(Qt.AlignCenter)
            header_item.setBackground(QBrush(QColor("#4a6fa5")))
            header_item.setForeground(QBrush(QColor("#ffffff")))
            table.setVerticalHeaderItem(row, header_item)
        
        for row in range(len(table_rows)):
            row_values = table_rows[row]
            year = years[row]
            for col, value in enumerate(row_values):
                item = QTableWidgetItem(value)
                item.setTextAlignment(Qt.AlignCenter)
                item.setForeground(QBrush(QColor("#000000")))
                
                if value and value != "":
                    item.setForeground(QColor("#28a745"))
                    item.setToolTip(f"{headers[col]} {year}")
                else:
                    item.setBackground(QColor("#9a9a73"))
                
                table.setItem(row, col, item)
                item.setData(Qt.UserRole + 2, True)
        
        table.resizeColumnsToContents()
        
        progress.close()
        
        QMessageBox.information(
            self,
            "Импорт завершён",
            f"✅ Таблица создана!\n\n"
            f"📊 Страна: {country_name}\n"
            f"📅 Всего строк: {len(table_rows)}\n"
            f"🪙 Номиналов: {len(headers)}\n"
            f"📄 Лист: {sheet_name}\n\n"
            f"Не забудьте сохранить изменения (SAVE ALL)."
        )
        
        self.is_modified = True
        self.refresh_color_legend()

    def _force_reload(self):
        """Принудительная перезагрузка из файла (сбрасывает кэш)"""
        country_id = self.current_country_id
        country_name = self.current_country_name
        
        if not country_name:
            return
        
        cache_key = country_id
        if cache_key in self.loaded_tables_cache:
            self.logger.info(f"🗑️ Принудительно удаляем кэш для {cache_key}")
            del self.loaded_tables_cache[cache_key]
        
        self._loaded_country_id = None
        self._loaded_country_name = None
        
        self._load_country_table()

    def save_tab_order(self):
        """Сохраняет порядок вкладок в настройки"""
        if not self.sheet_tabs:
            return
        
        tab_order = []
        for i in range(self.sheet_tabs.count()):
            tab_text = self.sheet_tabs.tabText(i)
            tab_order.append(tab_text)
        
        if self.current_country_id is not None:
            settings_key = f'tab_order_country_{self.current_country_id}'
        else:
            settings_key = f'tab_order_all'
        
        self.settings.setValue(settings_key, tab_order)
        self.logger.debug(f"💾 Сохранён порядок вкладок: {tab_order}")
    
    def restore_tab_order(self):
        """Восстанавливает порядок вкладок из настроек"""
        if not self.sheet_tabs:
            return
        
        if self.current_country_id is not None:
            settings_key = f'tab_order_country_{self.current_country_id}'
        else:
            settings_key = f'tab_order_all'
        
        saved_order = self.settings.value(settings_key)
        if not saved_order or not isinstance(saved_order, list):
            return
        
        current_tabs = []
        for i in range(self.sheet_tabs.count()):
            tab_text = self.sheet_tabs.tabText(i)
            current_tabs.append((i, tab_text))
        
        for target_pos, tab_name in enumerate(saved_order):
            for current_pos, current_name in current_tabs:
                if current_name == tab_name:
                    if current_pos != target_pos:
                        widget = self.sheet_tabs.widget(current_pos)
                        self.sheet_tabs.removeTab(current_pos)
                        self.sheet_tabs.insertTab(target_pos, widget, tab_name)
                    break
        
        self.logger.debug(f"🔄 Восстановлен порядок вкладок: {saved_order}")

    def load_dataframe(self, df):
        """Загружает данные из DataFrame"""
        try:
            import pandas as pd
        except ImportError:
            QMessageBox.warning(self, "Ошибка", "Библиотека pandas не установлена")
            return
        
        self.sheet_tabs.clear()
        self.sheets.clear()
        self.add_new_sheet("Лист1")
        
        sheet = self.sheets["Лист1"]
        table = sheet.table
        
        table.setRowCount(len(df))
        table.setColumnCount(len(df.columns))
        
        for i, col in enumerate(df.columns):
            header_item = QTableWidgetItem(str(col))
            header_item.setBackground(QBrush(QColor("#4a6fa5")))
            header_item.setForeground(QBrush(QColor("#ffffff")))
            header_item.setTextAlignment(Qt.AlignCenter)
            table.setHorizontalHeaderItem(i, header_item)
        
        for i, row in df.iterrows():
            for j, val in enumerate(row):
                item = QTableWidgetItem(str(val) if pd.notna(val) else "")
                item.setForeground(QBrush(QColor("#000000")))
                if pd.isna(val) or val == "":
                    item.setBackground(QBrush(QColor("#6c757d")))
                    item.setForeground(QBrush(QColor("#ffffff")))
                table.setItem(i, j, item)
                item.setData(Qt.UserRole + 2, True)
        
        self.update_size_info()
        self.refresh_color_legend()
        self.is_modified = True

    def closeEvent(self, event):
        """Обработчик закрытия"""
        if self.is_modified:
            self.save_current_table()
        super().closeEvent(event)
        
    def _get_sheets_order_file(self):
        """Путь к файлу с порядком вкладок"""
        from utils.paths import paths
        return paths.get_data_dir() / "yearly_tables" / "sheets_order.json"

    def _on_sheet_tab_moved(self):
        """Пользователь перетащил вкладку — отложенно сохраняем порядок"""
        from PySide6.QtCore import QTimer
        QTimer.singleShot(0, self._save_sheet_order)

    def _save_sheet_order(self):
        """Сохраняет текущий порядок вкладок погодовки"""
        try:
            import json
            order = [self.sheet_tabs.tabText(i) for i in range(self.sheet_tabs.count())]
            path = self._get_sheets_order_file()
            path.parent.mkdir(parents=True, exist_ok=True)
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(order, f, ensure_ascii=False)
            self.logger.info(f"💾 Порядок вкладок сохранён: {order}")
        except Exception as e:
            self.logger.error(f"Ошибка сохранения порядка вкладок: {e}")

    def _restore_sheet_order(self):
        """Восстанавливает сохранённый порядок вкладок"""
        try:
            import json
            path = self._get_sheets_order_file()
            if not path.exists():
                return
            with open(path, 'r', encoding='utf-8') as f:
                order = json.load(f)
            bar = self.sheet_tabs.tabBar()
            for target_pos, name in enumerate(order):
                current_pos = -1
                for i in range(self.sheet_tabs.count()):
                    if self.sheet_tabs.tabText(i) == name:
                        current_pos = i
                        break
                if current_pos > target_pos:
                    bar.moveTab(current_pos, target_pos)
        except Exception as e:
            self.logger.error(f"Ошибка восстановления порядка вкладок: {e}")
            
    def _save_all_with_order(self):
        """SAVE ALL: сначала порядок листов → в визуальный, затем оригинальное сохранение кэша"""
        self._sync_sheets_visual_order()
        self._orig_save_all_sheets()
        self._save_sheet_order()

    def _sync_sheets_visual_order(self):
        """Переставляет self.sheets в визуальный порядок вкладок, чтобы кэш сохранил его"""
        try:
            if not hasattr(self, 'sheets') or not isinstance(self.sheets, dict):
                return
            # карта: объект SheetWidget -> ключ листа
            widget_to_key = {}
            for key, sheet in self.sheets.items():
                widget_to_key[id(sheet)] = key
            ordered = {}
            for i in range(self.sheet_tabs.count()):
                w = self.sheet_tabs.widget(i)
                key = widget_to_key.get(id(w))
                if key is not None:
                    ordered[key] = self.sheets[key]
            # остальное (листы вне вкладок, если есть) — в конец
            for key, sheet in self.sheets.items():
                if key not in ordered:
                    ordered[key] = sheet
            self.sheets.clear()
            self.sheets.update(ordered)
            self.logger.info(f"🔀 Порядок листов синхронизирован с вкладками: {list(ordered.keys())}")
        except Exception as e:
            self.logger.error(f"Ошибка синхронизации порядка листов: {e}")
       
    def _show_sheet_context_menu(self, position):
        """Контекстное меню вкладки: переименовать / закрыть"""
        from PySide6.QtWidgets import QMenu, QInputDialog
        bar = self.sheet_tabs.tabBar()
        index = bar.tabAt(position)
        if index < 0:
            return
        menu = QMenu(self)
        rename_action = menu.addAction("✏️ Переименовать вкладку")
        menu.addSeparator()
        close_action = menu.addAction("❌ Закрыть вкладку")
        action = menu.exec(bar.mapToGlobal(position))
        if action == rename_action:
            old_name = self.sheet_tabs.tabText(index)
            new_name, ok = QInputDialog.getText(
                self, "Переименовать вкладку", "Новое название:", text=old_name
            )
            if ok and new_name.strip() and new_name.strip() != old_name:
                self.sheet_tabs.setTabText(index, new_name.strip())
                self._on_sheet_renamed(index, old_name, new_name.strip())
                self._save_sheet_order()
        elif action == close_action:
            self.close_sheet(index)

    def _on_sheet_renamed(self, index, old_name, new_name):
        """Обновляет имя листа во внутренних данных после переименования"""
        try:
            if hasattr(self, 'sheets') and isinstance(self.sheets, dict):
                if old_name in self.sheets:
                    sheet = self.sheets.pop(old_name)
                    sheet.sheet_name = new_name
                    self.sheets[new_name] = sheet
                else:
                    # fallback: ищем лист по виджету вкладки
                    w = self.sheet_tabs.widget(index)
                    for key, sheet in list(self.sheets.items()):
                        if sheet is w:
                            self.sheets.pop(key)
                            sheet.sheet_name = new_name
                            self.sheets[new_name] = sheet
                            break
                self.is_modified = True
            self.logger.info(f"✏️ Вкладка переименована: '{old_name}' → '{new_name}'")
        except Exception as e:
            self.logger.error(f"Ошибка при переименовании листа: {e}")