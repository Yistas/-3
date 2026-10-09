# -*- coding: utf-8 -*-

"""
Импорт/экспорт для вкладки "Сделки"
"""

import csv
import logging
from pathlib import Path

from PySide6.QtCore import QThread, Signal, QTimer
from PySide6.QtWidgets import QFileDialog, QMessageBox, QProgressBar

from database.models import Deal


class ImportDealsThread(QThread):
    """Поток для импорта сделок из Excel"""

    progress = Signal(int, int, str)
    finished = Signal(int, int)
    error = Signal(str)

    EXCEL_TO_DB = {
        'НИК': 'buyer',
        'Сумма': 'amount',
        'ДОСТ': 'delivery',
        'ИТОГ': 'total',
        'Рез': 'reserve',
        'Собр': 'collect',
        'СКАН': 'scan',
        'Отосл': 'send',
        'Согл': 'approve',
        'Оплата': 'payment',
        'Пров': 'check',
        'Упак': 'pack',
        'почта': 'post',
        'ДОШЛО': 'arrived',
        '№№': 'number',
        'тип ': 'type',
        'где': 'where'
    }

    def __init__(self, db_manager, file_path):
        super().__init__()
        self.db_manager = db_manager
        self.file_path = file_path
        self._is_running = True
        self.logger = logging.getLogger('CoinCollector.GUI.ImportDealsThread')

    def run(self):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(self.file_path, data_only=True)
            ws = wb.active

            headers = []
            for cell in ws[1]:
                headers.append(str(cell.value).strip() if cell.value else "")

            col_indices = {}
            for excel_col, db_col in self.EXCEL_TO_DB.items():
                for idx, header in enumerate(headers):
                    if header == excel_col:
                        col_indices[db_col] = idx
                        break

            imported = 0
            errors = 0

            try:
                self.db_manager.session.query(Deal).delete()
                self.db_manager.session.commit()
            except:
                pass

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not self._is_running:
                    break

                try:
                    deal = Deal()
                    has_data = False

                    for db_col, col_idx in col_indices.items():
                        if col_idx < len(row) and row[col_idx] is not None:
                            value = row[col_idx]
                            str_value = str(value).strip()

                            if str_value and str_value not in ['nan', 'None', '']:
                                if db_col == 'number':
                                    try:
                                        num_val = float(value)
                                        if num_val == int(num_val):
                                            str_value = str(int(num_val))
                                        else:
                                            str_value = str(num_val)
                                    except (ValueError, TypeError):
                                        pass

                                setattr(deal, db_col, str_value)
                                has_data = True

                    if has_data:
                        self.db_manager.session.add(deal)
                        imported += 1

                        if imported % 100 == 0:
                            self.db_manager.session.commit()

                except Exception as e:
                    errors += 1
                    self.logger.debug(f"Ошибка в строке {row_idx}: {e}")

            self.db_manager.session.commit()
            self.finished.emit(imported, errors)

        except Exception as e:
            import traceback
            traceback.print_exc()
            self.error.emit(str(e))

    def stop(self):
        self._is_running = False


class DealsIOMixin:
    """Примесь с методами импорта/экспорта"""

    def _import_from_excel(self):
        """Импорт сделок из Excel"""
        self.logger.debug("_import_from_excel: импорт из Excel")

        file_path, _ = QFileDialog.getOpenFileName(
            self, "Импорт сделок из Excel",
            "",
            "Excel files (*.xlsx *.xls);;All files (*.*)"
        )

        if not file_path:
            return

        reply = QMessageBox.question(
            self, "Подтверждение",
            f"Импортировать сделки из файла?\n\n{Path(file_path).name}",
            QMessageBox.Yes | QMessageBox.No
        )

        if reply != QMessageBox.Yes:
            return

        self.progress_bar = QProgressBar()
        self.progress_bar.setMaximumHeight(3)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.show()
        self.progress_bar.setRange(0, 0)
        self.status_label.setText("⏳ Импорт сделок...")

        self._import_thread = ImportDealsThread(self.db_manager, file_path)
        self._import_thread.finished.connect(self._on_import_finished)
        self._import_thread.error.connect(self._on_import_error)
        self._import_thread.start()

    def _on_import_finished(self, imported, errors):
        """Обработчик завершения импорта"""
        self.logger.debug(f"_on_import_finished: импортировано {imported}, ошибок {errors}")
        self.progress_bar.hide()

        try:
            self.db_manager.session.rollback()
        except:
            pass

        self.load_data()

        msg = f"✅ Импортировано: {imported}"
        if errors > 0:
            msg += f" | ❌ Ошибок: {errors}"

        self.status_label.setText(msg)
        self.status_label.setStyleSheet("color: #28a745; font-size: 11px;")
        QMessageBox.information(self, "Импорт завершён", msg)
        QTimer.singleShot(2000, lambda: self.status_label.setStyleSheet("color: #666; font-size: 11px;"))

    def _on_import_error(self, error_msg):
        """Обработчик ошибки импорта"""
        self.logger.error(f"_on_import_error: {error_msg}")
        self.progress_bar.hide()
        self.status_label.setText(f"❌ Ошибка: {error_msg[:50]}")
        self.status_label.setStyleSheet("color: #dc3545; font-size: 11px;")
        QMessageBox.critical(self, "Ошибка", f"Не удалось импортировать:\n{error_msg}")