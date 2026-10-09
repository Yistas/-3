# ===== gui/widgets/exchange_tab/purchases_io.py =====
# ПОЛНОСТЬЮ ЗАМЕНИТЬ ФАЙЛ

# -*- coding: utf-8 -*-

"""
Импорт/экспорт закупок (без pandas)
"""

import logging
from pathlib import Path
from PySide6.QtWidgets import QFileDialog, QMessageBox
from PySide6.QtCore import QThread, Signal

from openpyxl import load_workbook
from database.models import ImportCountry, Purchase


class ImportPurchasesThread(QThread):
    """Поток для импорта данных из Excel в БД (без pandas)"""
    progress = Signal(int, int, str)
    finished = Signal(int, int, int)
    error = Signal(str)
    
    def __init__(self, db_manager, file_path):
        super().__init__()
        self.db_manager = db_manager
        self.file_path = file_path
        self._is_running = True
    
    def run(self):
        try:
            # Загружаем Excel через openpyxl
            wb = load_workbook(self.file_path, data_only=True)
            
            # Ищем лист "ЗАКУПКИ"
            if "ЗАКУПКИ" in wb.sheetnames:
                ws = wb["ЗАКУПКИ"]
            else:
                ws = wb.active
            
            # Находим строку с заголовками
            header_row = None
            header_indices = {}
            
            rows = list(ws.iter_rows(values_only=True))
            
            for i, row in enumerate(rows):
                if not row:
                    continue
                row_values = [str(v).strip() if v else "" for v in row]
                if 'Старана' in row_values or 'ГОД' in row_values:
                    header_row = i
                    # Определяем индексы колонок
                    for col_idx, val in enumerate(row):
                        val_str = str(val).strip() if val else ""
                        if val_str in ['Старана', 'Страна']:
                            header_indices['country'] = col_idx
                        elif val_str == 'ГОД':
                            header_indices['year'] = col_idx
                        elif val_str == 'Номинал':
                            header_indices['denomination'] = col_idx
                        elif val_str == 'Валюта':
                            header_indices['currency'] = col_idx
                        elif val_str == 'Континент':
                            header_indices['continent'] = col_idx
                        elif val_str == 'КОЛ' or val_str == 'Кол-во':
                            header_indices['quantity'] = col_idx
                        elif val_str == 'КОМЕНТЫ':
                            header_indices['comments'] = col_idx
                    break
            
            if header_row is None:
                self.error.emit("Не найдена строка с заголовками (ищем 'Старана' или 'ГОД')")
                return
            
            # Обрабатываем строки данных
            total_rows = len(rows) - header_row - 1
            imported = 0
            skipped = 0
            errors = 0
            
            for idx in range(header_row + 1, len(rows)):
                if not self._is_running:
                    break
                
                row = rows[idx]
                if not row:
                    skipped += 1
                    continue
                
                try:
                    # Получаем страну
                    country_col = header_indices.get('country')
                    if country_col is None or country_col >= len(row):
                        skipped += 1
                        continue
                    
                    country_name = str(row[country_col]).strip() if row[country_col] else ""
                    if not country_name or country_name in ['nan', 'None', '']:
                        skipped += 1
                        continue
                    
                    # Получаем или создаём страну в таблице import_countries
                    import_country = self.db_manager.session.query(ImportCountry).filter_by(
                        name=country_name
                    ).first()
                    
                    if not import_country:
                        import_country = ImportCountry(name=country_name)
                        self.db_manager.session.add(import_country)
                        self.db_manager.session.flush()
                    
                    # Год
                    year = None
                    year_col = header_indices.get('year')
                    if year_col is not None and year_col < len(row) and row[year_col]:
                        try:
                            year = int(float(str(row[year_col])))
                        except:
                            pass
                    
                    # Номинал
                    denomination = None
                    denom_col = header_indices.get('denomination')
                    if denom_col is not None and denom_col < len(row) and row[denom_col]:
                        denomination = str(row[denom_col]).strip()
                        if denomination in ['nan', 'None', '']:
                            denomination = None
                    
                    # Валюта
                    currency = None
                    currency_col = header_indices.get('currency')
                    if currency_col is not None and currency_col < len(row) and row[currency_col]:
                        currency = str(row[currency_col]).strip()
                        if currency in ['nan', 'None', '']:
                            currency = None
                    
                    # Континент
                    continent = None
                    continent_col = header_indices.get('continent')
                    if continent_col is not None and continent_col < len(row) and row[continent_col]:
                        continent = str(row[continent_col]).strip()
                        if continent in ['nan', 'None', '']:
                            continent = None
                    
                    # Количество
                    quantity = 1
                    qty_col = header_indices.get('quantity')
                    if qty_col is not None and qty_col < len(row) and row[qty_col]:
                        try:
                            quantity = int(float(str(row[qty_col])))
                        except:
                            pass
                    
                    # Комментарии
                    comments = None
                    comments_col = header_indices.get('comments')
                    if comments_col is not None and comments_col < len(row) and row[comments_col]:
                        comments = str(row[comments_col]).strip()
                        if comments in ['nan', 'None', '']:
                            comments = None
                    
                    purchase = Purchase(
                        import_country_id=import_country.id,
                        continent=continent,
                        denomination_value=denomination,
                        currency=currency,
                        year=year,
                        quantity=quantity,
                        comments=comments
                    )
                    
                    self.db_manager.session.add(purchase)
                    imported += 1
                    
                    if imported % 100 == 0:
                        self.db_manager.session.commit()
                        self.progress.emit(imported, total_rows, f"Импортировано {imported}...")
                    
                except Exception as e:
                    errors += 1
                    self.progress.emit(imported + errors, total_rows, f"❌ Ошибка в строке {idx + 1}: {str(e)[:50]}")
            
            self.db_manager.session.commit()
            self.finished.emit(imported, skipped, errors)
            
        except Exception as e:
            self.error.emit(str(e))
    
    def stop(self):
        self._is_running = False